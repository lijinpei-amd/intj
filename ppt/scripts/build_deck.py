# pyright: standard
"""Builds the deck outlined in ppt/ppt.md, once per theme.

usage: python ppt/scripts/build_deck.py [--theme amd|light|all] [--template ~/blank.potx]
writes ppt/results/intj_launch_overhead_amd.pptx (on the ~/blank.potx theme) and
ppt/results/intj_launch_overhead_light.pptx (a light theme of its own)

Every number on a chart or in a sentence that has a CSV behind it is read from
ppt/data/; the rest cite their source in the slide's footer.
"""

import argparse
import copy
import csv
import io
import pathlib
import statistics
import zipfile
from dataclasses import dataclass, replace
from typing import Optional

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = pathlib.Path(__file__).resolve().parent.parent
DATA = ROOT / "data"

# ---- data -----------------------------------------------------------------


def read_csv(name):
    with open(DATA / name, newline="") as f:
        return list(csv.DictReader(f))


def f1(v):
    return f"{v:.1f}"


def f2(v):
    return f"{v:.2f}"


M308, M350, M4080 = "MI308X (gfx942)", "MI350X (gfx950)", "RTX 4080 SUPER (sm_89)"
MACHINES = [M308, M350, M4080]
SHORT = {M308: "MI308X", M350: "MI350X", M4080: "RTX 4080 SUPER"}

headline = read_csv("headline_launch.csv")
decode = read_csv("headline_decode_by_access_mode.csv")
sweep_args = read_csv("sweep_args_wide.csv")
sweep_const = read_csv("sweep_constexpr_wide.csv")
sweep_cache = read_csv("sweep_cache_wide.csv")
same_binary = read_csv("same_binary_1arg.csv")
ffi_sweep = read_csv("ffi_sweep.csv")
ffi_compare = read_csv("ffi_compare.csv")
kda = read_csv("aiter_kda_decode.csv")
flash = read_csv("aiter_flash_kda.csv")
fp8 = read_csv("aiter_fp8_mqa_logits.csv")
aiter_summary = read_csv("aiter_summary.csv")
fp8_proto = read_csv("aiter_fp8_mqa_logits_prototype.csv")
incompat = read_csv("incompatible_micro.csv")
kmap = read_csv("kernel_cache_map.csv")
ffi_disp = read_csv("ffi_dispatch.csv")
fp8_breakdown = read_csv("aiter_fp8_mqa_logits_breakdown.csv")


def hl(machine, grid, col):
    return float(
        next(r for r in headline if r["machine"] == machine and r["grid"] == grid)[col]
    )


def dec(machine, mode):
    row = next(
        r for r in decode if r["machine"] == machine and r["torch_access_mode"] == mode
    )
    return float(row["decode_spec_key_ns"])


def op_host(rows, launch_path, keep=lambda _r: True):
    return statistics.median(
        float(r["host_us"])
        for r in rows
        if r["route"] == "default" and r["launch_path"] == launch_path and keep(r)
    )


# The one name the deck uses for decoding a call's arguments into its spec key (AGENTS.md).
ARGS_KEY = "intj args -> spec-key"


# ---- themes -----------------------------------------------------------------

W = 13.333
H = 7.5
MX = 0.64  # the AMD template's title placeholder x
CW = 12.05
CODE = "Courier New"


@dataclass(frozen=True)
class Theme:
    name: str
    template: Optional[str]  # a .potx whose master and layouts the deck uses, or None
    font: str  # body text
    head_font: str  # slide titles
    bg: str  # slide background
    ink: str  # primary text
    ink2: str  # secondary text
    muted: str  # captions, axis labels
    card: str  # card body, odd table rows
    card2: str  # card header, table header
    head_ink: str  # table header text
    mid: str  # neutral badge fill
    on_mid: str  # text on a neutral badge
    on_accent: str  # text on an intj-colored fill
    on_series: str  # data labels inside bars
    rule: str  # neutral card rule
    code_bg: str
    code_line: str
    code_ink: str
    grid: str  # chart gridlines
    axis: str  # chart category axis
    triton: str  # upstream / baseline series
    intj: str  # intj series and the deck accent
    third: str  # TVM FFI, AITER _FastLaunch, GPU kernel time
    driver: str  # driver share of a Triton launch


# The AMD corporate template's palette (its "Corporate Presentation Color Palette" slide).
AMD = Theme(
    name="amd",
    template=str(pathlib.Path.home() / "blank.potx"),
    font="Arial",
    head_font="Arial",
    bg="000000",
    ink="FFFFFF",
    ink2="D9D9D9",
    muted="A6A6A6",
    card="262626",
    card2="404040",
    head_ink="FFFFFF",
    mid="595959",
    on_mid="FFFFFF",
    on_accent="000000",
    on_series="FFFFFF",
    rule="C1A968",
    code_bg="141414",
    code_line="404040",
    code_ink="D9D9D9",
    grid="333333",
    axis="595959",
    triton="808080",
    intj="00C2DE",
    third="C1A968",
    driver="404040",
)

# A light theme of its own; the three series colors pass the dataviz palette validator on white.
LIGHT = Theme(
    name="light",
    template=None,
    font="Calibri",
    head_font="Arial",
    bg="FFFFFF",
    ink="1B1F24",
    ink2="4A4F57",
    muted="7A7F87",
    card="F1F3F6",
    card2="E1E5EA",
    head_ink="1B1F24",
    mid="6B7079",
    on_mid="FFFFFF",
    on_accent="1B1F24",
    on_series="FFFFFF",
    rule="6B7079",
    code_bg="1F2329",
    code_line="1F2329",
    code_ink="E3E6EA",
    grid="E7E9ED",
    axis="C9CDD3",
    triton="2A78D6",
    intj="EB6834",
    third="1BAF7A",
    driver="6B7079",
)

T = AMD  # the theme being built; set by build()


def rgb(hex_):
    return RGBColor.from_string(hex_)


def open_template(path):
    """Opens a .potx as a presentation by retyping its main part."""
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(
        buf, "w", zipfile.ZIP_DEFLATED
    ) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"presentationml.template.main+xml",
                    b"presentationml.presentation.main+xml",
                )
            zout.writestr(item, data)
    buf.seek(0)
    prs = Presentation(buf)
    sld_ids = prs.slides._sldIdLst  # pyright: ignore[reportPrivateUsage] -- no public slide delete
    for sld_id in list(sld_ids):
        prs.part.drop_rel(sld_id.rId)
        sld_ids.remove(sld_id)
    return prs


# ---- primitives -------------------------------------------------------------


def _run_style(run, size, color, bold, italic, font):
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def _sub(parent, tag):
    el = OxmlElement(tag)
    parent.append(el)
    return el


def _bullet(paragraph, color):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.22)))
    pPr.set("indent", str(-Inches(0.22)))
    clr = _sub(pPr, "a:buClr")
    _sub(clr, "a:srgbClr").set("val", color)
    _sub(pPr, "a:buFont").set("typeface", "Arial")
    _sub(pPr, "a:buChar").set("char", "•")


def _no_bullet(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    _sub(pPr, "a:buNone")


def fill_frame(
    tf,
    paras,
    size=14,
    color=None,
    bold=False,
    italic=False,
    font=None,
    align="l",
    anchor="t",
    bullets=False,
    gap=0,
):
    """Writes paragraphs into a text frame.

    paras: a str ("\n" splits paragraphs), or a list whose items are a str or a
    list of runs; a run is a str or (text, {size, color, bold, italic, font}).
    """
    color = color or T.ink
    font = font or T.font
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {
        "t": MSO_ANCHOR.TOP,
        "m": MSO_ANCHOR.MIDDLE,
        "b": MSO_ANCHOR.BOTTOM,
    }[anchor]
    if isinstance(paras, str):
        paras = paras.split("\n")
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[
            align
        ]
        if gap and i < len(paras) - 1:
            p.space_after = Pt(gap)
        if bullets:
            _bullet(p, T.muted)
        else:
            _no_bullet(p)
        runs = [para] if isinstance(para, (str, tuple)) else para
        for r in runs:
            txt, o = (r, {}) if isinstance(r, str) else r
            run = p.add_run()
            run.text = txt
            _run_style(
                run,
                o.get("size", size),
                o.get("color", color),
                o.get("bold", bold),
                o.get("italic", italic),
                o.get("font", font),
            )


def text(slide, x, y, w, h, paras, **kw):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_frame(box.text_frame, paras, **kw)
    return box


def B(t, **o):
    return (t, {"bold": True, **o})


def C_(t, **o):
    return (t, {"font": CODE, **o})


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    # The shape's theme style reference brings a shadow in some themes; every fill and line is explicit.
    style = shp._element.find(qn("p:style"))  # pyright: ignore[reportPrivateUsage]
    if style is not None:
        shp._element.remove(style)  # pyright: ignore[reportPrivateUsage]
    return shp


def badge(slide, x, y, glyph, fill=None, d=0.42, ink=None, size=14):
    fill, ink = fill or T.intj, ink or T.on_accent
    shp = rect(slide, x, y, d, d, fill, shape=MSO_SHAPE.OVAL)
    fill_frame(
        shp.text_frame, glyph, size=size, color=ink, bold=True, align="c", anchor="m"
    )
    return shp


def arrow(slide, x, y1, y2, color=None):
    color = color or T.muted
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y1), Inches(x), Inches(y2)
    )
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(1.5)
    ln = conn.line._get_or_add_ln()  # pyright: ignore[reportPrivateUsage] -- no public arrowhead API
    _sub(ln, "a:tailEnd").set("type", "triangle")
    return conn


def card(slide, x, y, w, h, fill=None):
    return rect(slide, x, y, w, h, fill or T.card)


def header_card(slide, x, y, w, h, title, rule=None, head_h=0.5, size=15):
    """The template's text box: a colored rule, a Dark Gray 2 header, a Dark Gray 1 body."""
    rect(slide, x, y, w, 0.05, rule or T.rule)
    rect(slide, x, y + 0.05, w, head_h, T.card2)
    rect(slide, x, y + 0.05 + head_h, w, h - 0.05 - head_h, T.card)
    text(
        slide,
        x + 0.2,
        y + 0.05,
        w - 0.4,
        head_h,
        title,
        size=size,
        bold=True,
        anchor="m",
    )
    return y + 0.05 + head_h  # top of the body


def stat(slide, x, y, w, h, big, label, color=None, big_size=28, label_size=12):
    color = color or T.intj
    card(slide, x, y, w, h)
    text(
        slide,
        x + 0.2,
        y + 0.12,
        w - 0.4,
        h * 0.5,
        big,
        size=big_size,
        bold=True,
        color=color,
        anchor="m",
    )
    text(
        slide,
        x + 0.2,
        y + 0.12 + h * 0.5,
        w - 0.4,
        h * 0.5 - 0.2,
        label,
        size=label_size,
        color=T.ink2,
    )


def code_box(slide, x, y, w, h, code, size=12):
    rect(slide, x, y, w, h, T.code_bg, line=T.code_line)
    text(
        slide,
        x + 0.22,
        y + 0.12,
        w - 0.44,
        h - 0.24,
        code,
        size=size,
        font=CODE,
        color=T.code_ink,
        anchor="m",
    )


DIFF_ADD, DIFF_DEL = "7EE787", "FF7B72"  # on the dark code background of both themes


def diff_box(slide, x, y, w, h, lines, size=11):
    """A code box that colors unified-diff lines by their first character."""
    rect(slide, x, y, w, h, T.code_bg, line=T.code_line)
    paras = []
    for ln in lines:
        color = (
            DIFF_ADD
            if ln.startswith("+")
            else DIFF_DEL
            if ln.startswith("-")
            else T.code_ink
        )
        paras.append([(ln or " ", {"color": color})])
    text(slide, x + 0.22, y + 0.15, w - 0.44, h - 0.3, paras, size=size, font=CODE)


def picture(slide, x, y, w, image, pad=0.12):
    """A figure on a white mat, w wide; returns the bottom of the mat."""
    if isinstance(image, Image.Image):
        buf = io.BytesIO()
        image.save(buf, format="PNG")
        buf.seek(0)
        size, src = image.size, buf
    else:
        with Image.open(image) as im:
            size = im.size
        src = str(image)
    h = (w - 2 * pad) * size[1] / size[0] + 2 * pad
    rect(slide, x, y, w, h, "FFFFFF", line=None if T.bg == "000000" else T.card2)
    slide.shapes.add_picture(
        src, Inches(x + pad), Inches(y + pad), width=Inches(w - 2 * pad)
    )
    return y + h


def legend(slide, x, y, items):
    """items: (name, color) or (name, color, "dash") for a dashed line series."""
    cx = x
    for name, color, *style in items:
        if style == ["dash"]:
            conn = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(cx),
                Inches(y + 0.16),
                Inches(cx + 0.3),
                Inches(y + 0.16),
            )
            conn.line.color.rgb = rgb(color)
            conn.line.width = Pt(2.25)
            conn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            cx += 0.13
        else:
            rect(slide, cx, y + 0.08, 0.17, 0.17, color)
        w = 0.25 + len(name) * 0.085
        text(slide, cx + 0.26, y, w, 0.32, name, size=11, color=T.ink2, anchor="m")
        cx += 0.26 + w + 0.2


def table(slide, x, y, w, rows, col_w, row_h=0.42, size=12):
    """rows[0] is the header; a cell is a str or (text, {color, bold, font})."""
    shape = slide.shapes.add_table(
        len(rows),
        len(rows[0]),
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(row_h * len(rows)),
    )
    tbl = shape.table
    # "No Style, No Grid": every fill below is explicit.
    style_id = tbl._tbl.tblPr.find(qn("a:tableStyleId"))  # pyright: ignore[reportPrivateUsage]
    if style_id is None:
        style_id = _sub(tbl._tbl.tblPr, "a:tableStyleId")  # pyright: ignore[reportPrivateUsage]
    style_id.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for j, cw_ in enumerate(col_w):
        tbl.columns[j].width = Inches(cw_)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(
                T.card2 if i == 0 else (T.card if i % 2 else T.bg)
            )
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            txt, o = (val, {}) if isinstance(val, str) else val
            fill_frame(
                cell.text_frame,
                txt,
                size=o.get("size", size),
                color=o.get("color", T.head_ink if i == 0 else T.ink),
                bold=o.get("bold", i == 0),
                font=o.get("font", T.font),
                anchor="m",
                align=o.get("align", "l"),
            )
    return tbl


def source(slide, s):
    text(slide, MX, 6.93, 10.9, 0.25, s, size=8, color=T.muted, anchor="m")


# ---- charts ---------------------------------------------------------------


def _transparent(chart):
    cs = chart._chartSpace  # pyright: ignore[reportPrivateUsage] -- chart-area fill has no API
    sp = parse_xml(
        '<c:spPr xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        "<a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>"
    )
    cs.find(qn("c:chart")).addnext(sp)
    plot_area = cs.find(qn("c:chart")).find(qn("c:plotArea"))
    ext = plot_area.find(qn("c:extLst"))
    if ext is not None:
        ext.addprevious(copy.deepcopy(sp))
    else:
        plot_area.append(copy.deepcopy(sp))


def _style_axes(chart, y_title=None, x_title=None, y_max=None):
    chart.has_legend = False
    chart.has_title = False
    chart.font.name = T.font
    chart.font.size = Pt(11)
    chart.font.color.rgb = rgb(T.ink2)
    va, ca = chart.value_axis, chart.category_axis
    va.minimum_scale = 0
    if y_max is not None:
        va.maximum_scale = y_max
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(T.grid)
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(10)
    va.tick_labels.font.color.rgb = rgb(T.muted)
    ca.has_major_gridlines = False
    ca.format.line.color.rgb = rgb(T.axis)
    ca.tick_labels.font.size = Pt(11)
    ca.tick_labels.font.color.rgb = rgb(T.ink2)
    for axis, title in ((va, y_title), (ca, x_title)):
        if title:
            axis.has_title = True
            fill_frame(axis.axis_title.text_frame, title, size=10, color=T.muted)
    _transparent(chart)


def bar_chart(
    slide,
    x,
    y,
    w,
    h,
    cats,
    series,
    horizontal=False,
    stacked=False,
    fmt="0.0",
    label_pos=XL_LABEL_POSITION.OUTSIDE_END,
    label_color=None,
    gap=60,
    label_size=11,
    y_max=None,
):
    data = CategoryChartData()
    data.categories = cats
    for name, values, _ in series:
        data.add_series(name, values)
    kind = {
        (False, False): XL_CHART_TYPE.COLUMN_CLUSTERED,
        (False, True): XL_CHART_TYPE.COLUMN_STACKED,
        (True, False): XL_CHART_TYPE.BAR_CLUSTERED,
        (True, True): XL_CHART_TYPE.BAR_STACKED,
    }[(horizontal, stacked)]
    chart = slide.shapes.add_chart(
        kind, Inches(x), Inches(y), Inches(w), Inches(h), data
    ).chart
    _style_axes(chart, y_max=y_max)
    plot = chart.plots[0]
    plot.gap_width = gap
    plot.overlap = 100 if stacked else -10
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = fmt
    labels.number_format_is_linked = False
    labels.position = label_pos
    labels.font.size = Pt(label_size)
    labels.font.color.rgb = rgb(label_color or T.ink)
    labels.font.bold = True
    for s, (_, _, color) in zip(plot.series, series):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(color)
        s.invert_if_negative = False
    return chart


def line_chart(slide, x, y, w, h, cats, series, x_title=None, y_max=None, y_step=None):
    data = CategoryChartData()
    data.categories = cats
    for name, values, *_ in series:
        data.add_series(name, values)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), data
    ).chart
    _style_axes(chart, x_title=x_title, y_max=y_max)
    if y_step:
        chart.value_axis.major_unit = y_step
    for s, (_, _, color, *style) in zip(chart.plots[0].series, series):
        s.smooth = False
        s.format.line.color.rgb = rgb(color)
        s.format.line.width = Pt(2.25)
        if style == ["dash"]:
            s.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        s.marker.style = XL_MARKER_STYLE.CIRCLE
        s.marker.size = 7
        s.marker.format.fill.solid()
        s.marker.format.fill.fore_color.rgb = rgb(color)
        s.marker.format.line.color.rgb = rgb(color)
    return chart


# ---- slides ---------------------------------------------------------------


class Deck:
    def __init__(self, theme):
        self.theme = theme
        if theme.template:
            self.prs = open_template(theme.template)
        else:
            self.prs = Presentation()
            self.prs.slide_width = Inches(W)
            self.prs.slide_height = Inches(H)
        self.layouts = {layout.name: layout for layout in self.prs.slide_layouts}

    def slide(self, title, layout="Title Only", notes=None):
        """Adds a slide; without a template the layout's title is restyled and a page number added."""
        s = self.prs.slides.add_slide(self.layouts[layout])
        if title is not None:
            assert s.shapes.title is not None
            # Longer titles wrap at 28 pt and run into the content below.
            assert layout != "Title Only" or len(title) <= 64, title
            s.shapes.title.text = title
        if not self.theme.template:
            s.background.fill.solid()
            s.background.fill.fore_color.rgb = rgb(self.theme.bg)
            if layout == "Title Only" and s.shapes.title is not None:
                t = s.shapes.title
                t.left, t.top, t.width, t.height = (
                    Inches(MX),
                    Inches(0.55),
                    Inches(CW),
                    Inches(0.6),
                )
                fill_frame(
                    t.text_frame, title, size=28, bold=True, font=self.theme.head_font
                )
                text(
                    s,
                    W - MX - 0.6,
                    6.93,
                    0.6,
                    0.25,
                    str(len(self.prs.slides)),
                    size=10,
                    color=self.theme.muted,
                    align="r",
                    anchor="m",
                )
        if notes:
            frame = s.notes_slide.notes_text_frame
            assert frame is not None
            frame.text = notes
        return s

    def placeholder(self, s, idx, paras):
        ph = next(p for p in s.placeholders if p.placeholder_format.idx == idx)
        tf = ph.text_frame
        tf.text = ""
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = para
        return ph


TITLE = "Optimizing Triton Kernel Launch Overhead"
SUBTITLE = "INTJ, an INcompatible Triton JIT launcher"
AUTHOR = "Li Jinpei · 2026-10-15"
AGENDA = [
    "Motivation",
    "Method",
    "Results: performance and AITER adoption",
    "Implementation details",
    "Discussion",
]


def build(theme):
    global T
    T = theme
    d = Deck(theme)
    if theme.template:
        s = d.slide(TITLE, "Title Slide - No Image")
        d.placeholder(s, 12, [SUBTITLE])
        text(s, 5.85, 4.8, 6.83, 0.35, AUTHOR, size=14, color=T.muted)
        s = d.slide("Agenda", "Agenda")
        d.placeholder(s, 10, AGENDA)
    else:
        light_title(d)
        light_agenda(d)

    motivation(d)
    method(d)
    usage_example(d)
    results_headline(d)
    results_sweeps(d)
    citing_tilelang(d)
    citing_tilelang_figures(d)
    results_ffi(d)
    results_aiter(d)
    results_aiter_crossover(d)
    results_adoption(d)
    impl_launch_syntax(d)
    impl_ptr_range(d)
    impl_lazy(d)
    impl_autotune(d)
    impl_internals(d)
    impl_tests(d)
    impl_micro(d)
    impl_ffi_dispatch(d)
    impl_annotation(d)
    discussion(d)
    if theme.template:
        d.slide(None, "Closing logo slide")
    return d.prs


def light_title(d):
    s = d.slide(None, "Blank")
    rect(s, 0, 0, W, H, "F3F4F6")
    rect(s, 7.9, 0, W - 7.9, H, T.intj)
    text(s, MX + 0.1, 2.3, 7.0, 1.9, TITLE, size=40, bold=True, font=T.head_font)
    text(s, MX + 0.1, 4.3, 7.0, 0.45, SUBTITLE, size=20, color=T.ink2)
    text(s, MX + 0.1, 4.9, 7.0, 0.35, AUTHOR, size=14, color=T.muted)


def light_agenda(d):
    s = d.slide("Agenda")
    for i, item in enumerate(AGENDA):
        y = 1.75 + i * 0.95
        badge(s, MX, y, str(i + 1), d=0.55, size=18)
        text(s, MX + 0.85, y, 9, 0.55, item, size=22, anchor="m")
        if i < len(AGENDA) - 1:
            rect(s, MX + 0.85, y + 0.75, 8, 0.01, T.card2)


def motivation(d):
    s = d.slide(
        "Most of a Triton launch is Python, before the driver",
        notes=(
            "Host share = Triton grid=(0,) / grid=(1,): Triton binds, keys and looks up, then skips the driver "
            "call on the empty grid. The 'long tail' observation is from the intj README: studying JITFunction "
            "showed no low-hanging fruit."
        ),
    )
    host = [hl(m, "grid=(0,)", "triton_us") for m in MACHINES]
    full = [hl(m, "grid=(1,)", "triton_us") for m in MACHINES]
    share = [round(100 * h / f) for h, f in zip(host, full)]
    args = [r for r in sweep_args if r["machine"] == M308 and r["arg_kind"] == "tensor"]
    a0, a64 = float(args[0]["triton_host_us"]), float(args[-1]["triton_host_us"])
    k_gpu = float(
        next(
            r for r in aiter_summary if r["op"] == "kda_decode" and r["shape"] == "1x12"
        )["gpu_kernel_us"]
    )
    k_eager = float(
        next(
            r
            for r in kda
            if r["batch"] == "1" and r["launch_path"] == "base_triton_jit"
        )["eager_us"]
    )

    text(
        s,
        MX,
        1.6,
        5.6,
        4.9,
        [
            [
                B("Host path dominates. "),
                f"kernel[grid](*args) binds, keys, looks up and checks in Python: "
                f"{min(share)}–{max(share)}% of a launch on all three machines.",
            ],
            [
                B("It scales with arguments. "),
                f"0 → 64 tensor arguments: {f1(a0)} → {f1(a64)} µs of host time per launch (MI308X).",
            ],
            [
                B("No single hot spot. "),
                "The cost is a long tail spread across the call, not one slow function.",
            ],
            [
                B("Small kernels turn host-bound. "),
                f"AITER's fused KDA decode runs {f1(k_gpu)} µs on the GPU, "
                f"yet each eager call takes {f1(k_eager)} µs (MI350X).",
            ],
        ],
        size=15,
        bullets=True,
        gap=14,
    )

    text(
        s, 6.75, 1.55, 5.9, 0.35, "Triton launch time per call (µs)", size=14, bold=True
    )
    legend(
        s,
        6.75,
        1.95,
        [("host path (grid=(0,))", T.triton), ("driver launch + rest", T.driver)],
    )
    rev = [2, 1, 0]  # a bar chart draws its first category at the bottom
    bar_chart(
        s,
        6.6,
        2.35,
        6.1,
        4.3,
        [SHORT[MACHINES[i]] for i in rev],
        [
            ("host path", [host[i] for i in rev], T.triton),
            ("driver launch + rest", [full[i] - host[i] for i in rev], T.driver),
        ],
        horizontal=True,
        stacked=True,
        label_pos=XL_LABEL_POSITION.CENTER,
        label_color=T.on_series,
        gap=45,
        label_size=12,
    )
    source(
        s,
        "Source: data/headline_launch.csv (grid=(0,) vs grid=(1,)); data/sweep_args_wide.csv; "
        "data/aiter_kda_decode.csv, aiter_summary.csv (batch 1).",
    )


def method(d):
    s = d.slide(
        "Re-implement the launch path in C, rendered per kernel",
        notes=(
            "A key finer than Triton's costs a redundant compile; a coarser one launches the wrong binary silently. "
            "test_spec_key_is_never_coarser_than_triton guards it, and a mutation test shows it can fail."
        ),
    )
    text(
        s,
        MX,
        1.55,
        5.7,
        0.35,
        "What JITFunction does on every call",
        size=14,
        bold=True,
    )
    code_box(
        s,
        MX,
        1.95,
        5.75,
        1.85,
        "key    = hash(source, args, options, knobs)\n"
        "kernel = cache.get_or_create(key,\n"
        "           lambda: compile(source, args, ...))\n"
        "grid   = grid_func(args)\n"
        "launch(kernel, device, stream, grid, args)",
        size=12,
    )
    for i, (n, head, body) in enumerate(
        [
            (
                "1",
                "Delegate to C as soon as possible",
                "The Python call lands directly in a rendered METH_FASTCALL C entry, specialized to the kernel's signature.",
            ),
            (
                "2",
                "Change the interface where it costs time",
                "Device, stream and grid are passed in; features that need per-call Python are refused. Hence INcompatible.",
            ),
        ]
    ):
        y = 4.1 + i * 1.3
        badge(s, MX, y, n)
        text(s, MX + 0.6, y + 0.02, 5.1, 0.35, head, size=15, bold=True)
        text(s, MX + 0.6, y + 0.42, 5.15, 0.85, body, size=13, color=T.ink2)

    X = 6.9
    DW = W - MX - X

    def step(y, h, head, body):
        card(s, X, y, DW, h)
        text(s, X + 0.25, y + 0.12, DW - 0.5, 0.32, head, size=14, bold=True)
        text(s, X + 0.25, y + 0.47, DW - 0.5, h - 0.5, body, size=12, color=T.ink2)

    def hop(y, label):
        arrow(s, X + 0.6, y, y + 0.36)
        text(s, X + 0.78, y + 0.03, 3, 0.3, label, size=11, italic=True, color=T.muted)

    step(
        1.55,
        0.9,
        "@make_launcher",
        "Records the kernel and options. No GPU query, no torch import.",
    )
    hop(2.47, "first call")
    step(
        2.87,
        0.9,
        "Render · compile · load",
        "entry.c.jinja → .so, then swap the launcher's ml_meth to the entry.",
    )
    hop(3.79, "every later call")
    rect(s, X, 4.19, DW, 1.65, T.card, line=T.intj)
    text(s, X + 0.25, 4.3, 3, 0.32, "Native entry (C)", size=14, bold=True)
    chips = [ARGS_KEY, "cache lookup", "grid", "driver launch"]
    first_w = 1.95
    cw = (DW - 0.5 - first_w - 0.1 * (len(chips) - 1)) / (len(chips) - 1)
    cx = X + 0.25
    for i, c in enumerate(chips):
        w = first_w if i == 0 else cw
        last = i == len(chips) - 1
        shp = rect(s, cx, 4.75, w, 0.5, T.mid if last else T.intj)
        cx += w + 0.1
        fill_frame(
            shp.text_frame,
            c,
            size=10,
            bold=True,
            color=T.on_mid if last else T.on_accent,
            align="c",
            anchor="m",
        )
    text(
        s,
        X + 0.25,
        5.38,
        DW - 0.5,
        0.3,
        "miss → Python callback → triton.compile → insert",
        size=11,
        italic=True,
        color=T.muted,
    )
    text(
        s,
        X,
        6.1,
        DW,
        0.35,
        [
            [
                B("Invariant: ", color=T.intj),
                "same intj spec key ⟹ same Triton specialization",
            ]
        ],
        size=13,
    )
    source(
        s,
        "Source: README.md (How does it work?), AGENTS.md (The invariant), intj/render/entry.c.jinja.",
    )


def usage_example(d):
    s = d.slide(
        "Usage example: AITER's fused KDA decode on intj",
        notes=(
            "Diff condensed from aiter 874d253dd3..5d710e670: '...' marks elided arguments, and dev, stream and grid "
            "sit on separate lines in the original. Kernel: aiter/ops/triton/_triton_kernels/gated_delta_rule/decode/"
            "fused_conv_recurrent_norm.py. Call site: aiter/ops/triton/gated_delta_net/fused_kda_decode.py. "
            "_decode_num_warps() returns 2 on gfx942, else 4. current_device_stream() is torch.cuda.current_device() "
            "and its raw stream. Numbers: data/aiter_kda_decode.csv and aiter_summary.csv, MI350X."
        ),
    )
    pw = (CW - 0.3) / 2
    panels = [
        (
            "Kernel: decorate once",
            "_triton_kernels/.../fused_conv_recurrent_norm.py",
            [
                " import triton",
                " import triton.language as tl",
                "+from intj import make_launcher",
                "",
                '+@make_launcher(dynamic_options=("num_warps",))',
                " @triton.jit",
                " def fused_conv_recurrent_norm_kernel(",
                "     x_ptr, conv_weight_ptr, conv_state_ptr, ...,",
                "     T: tl.int64,",
                "     H: tl.constexpr, K: tl.constexpr, ...,",
                "     stride_x_tok: tl.int64, ...,",
                " ):",
            ],
        ),
        (
            "Call site: device, stream, grid, then positional",
            "gated_delta_net/fused_kda_decode.py",
            [
                " grid = (batch, H)",
                "-fused_conv_recurrent_norm_kernel[grid](",
                "+dev, stream = current_device_stream()",
                "+fused_conv_recurrent_norm_kernel(",
                "+    dev, stream, grid,",
                "+    _decode_num_warps(),  # dynamic option",
                "     mixed_qkv, conv_weight, conv_state, ...,",
                "-    H=H, K=K, V=V, W=W,",
                "-    stride_x_tok=mixed_qkv.stride(0), ...,",
                '-    num_warps=2 if get_arch() == "gfx942" else 4,',
                "+    H, K, V, W,",
                "+    mixed_qkv.stride(0), ...,  # stride_x_tok",
                " )",
            ],
        ),
    ]
    for i, (head, path, lines) in enumerate(panels):
        x = MX + i * (pw + 0.3)
        text(s, x, 1.5, pw, 0.35, head, size=14, bold=True)
        text(s, x, 1.85, pw, 0.28, path, size=10, font=CODE, color=T.muted)
        diff_box(s, x, 2.2, pw, 2.95, lines)

    up, ij = op_host(kda, "base_triton_jit"), op_host(kda, "intj")
    eager = [
        float(r["eager_speedup_vs_upstream"])
        for r in aiter_summary
        if r["op"] == "kda_decode" and r["eager_speedup_vs_upstream"]
    ][:4]  # batch 1-16
    cw = (CW - 0.9) / 4
    for i, (head, body) in enumerate(
        [
            (
                "@make_launcher",
                "Records the kernel; the C entry is built on the first call.",
            ),
            (
                "dynamic_options",
                "num_warps is passed per call, right after the grid, and keyed.",
            ),
            (
                "current_device_stream()",
                "AITER's helper: the current device and its raw stream.",
            ),
        ]
    ):
        x = MX + i * (cw + 0.3)
        card(s, x, 5.4, cw, 1.4)
        text(
            s,
            x + 0.2,
            5.52,
            cw - 0.4,
            0.32,
            head,
            size=12,
            bold=True,
            font=CODE,
            color=T.intj,
        )
        text(s, x + 0.2, 5.9, cw - 0.4, 0.85, body, size=12, color=T.ink2)
    stat(
        s,
        MX + 3 * (cw + 0.3),
        5.4,
        cw,
        1.4,
        f"{f1(up / ij)}×",
        f"host time {f1(up)} → {f1(ij)} µs; eager {min(eager):.1f}–{max(eager):.1f}× (MI350X)",
        big_size=24,
        label_size=11,
    )
    source(
        s,
        "Source: aiter 874d253dd3..5d710e670 (condensed diff); data/aiter_kda_decode.csv, "
        "data/aiter_summary.csv. Outputs bitwise equal on every shape.",
    )


def sweep_row(machine, n_args, kind="tensor"):
    return next(
        r
        for r in sweep_args
        if r["machine"] == machine
        and r["n_args"] == str(n_args)
        and r["arg_kind"] == kind
    )


def results_headline(d):
    # Host path = everything but the driver, on one kernel for both sides: Triton on an empty grid,
    # intj as a no_gpu launcher (call, intj args -> spec-key, cache lookup).
    tr_host = [float(sweep_row(m, 4)["triton_host_us"]) for m in MACHINES]
    ij_host = [float(sweep_row(m, 4)["intj_host_us"]) for m in MACHINES]
    ij_enter = [float(sweep_row(m, 0)["intj_host_us"]) for m in MACHINES]
    interp_extra = [
        (dec(m, "interpreter") - dec(m, "static_compile")) / 1000 for m in MACHINES
    ]
    s = d.slide(
        f"3.5–6.2× faster launches; intj's own host time is ~{max(ij_host):.2f} µs",
        notes=(
            "Machines: MI308X (Xeon 8480C), MI350X (EPYC 9965), RTX 4080 SUPER (Ryzen 9 7950X3D); each pinned to "
            "one core. gfx950 hipModuleLaunchKernel flips between ~2.3 and ~3.1 us, so MI350X launch rows move by "
            "+-0.7 us. Bar chart kernel: bench_launch.noop, 3 tensors + 2 scalars + 1 constexpr. Host-path table: "
            "the sweep's empty kernel with 4 tensor arguments. Triton's host path is a launch on grid=(0,), which binds, "
            "keys and looks up, then skips the driver. intj's is a no_gpu launcher on CPU tensors: the real entry with "
            "the driver call compiled out. intj on grid=(0,) is not used here because it returns before decoding any "
            "argument; _C_spec_key is not used because it also builds a Python (bytes, int) result. "
            f"Entering C alone (0 arguments): {min(ij_enter):.3f}-{max(ij_enter):.3f} us. Reading tensors through "
            f"the interpreter (INTERPRETER mode) adds {min(interp_extra):.2f}-{max(interp_extra):.2f} us to "
            f"{ARGS_KEY}."
        ),
    )
    text(
        s,
        MX,
        1.55,
        7,
        0.35,
        "Per launch incl. driver, grid=(1,), µs (lower is better)",
        size=14,
        bold=True,
    )
    legend(s, MX, 1.95, [("Triton JITFunction", T.triton), ("intj", T.intj)])
    bar_chart(
        s,
        MX - 0.1,
        2.35,
        7.0,
        4.4,
        [SHORT[m] for m in MACHINES],
        [
            ("Triton", [hl(m, "grid=(1,)", "triton_us") for m in MACHINES], T.triton),
            ("intj", [hl(m, "grid=(1,)", "intj_us") for m in MACHINES], T.intj),
        ],
        fmt="0.00",
        gap=70,
        label_size=12,
    )

    X = 7.85
    TW = W - MX - X
    text(
        s,
        X,
        1.55,
        TW,
        0.35,
        "Host path, no driver, 4 tensor args (µs)",
        size=14,
        bold=True,
    )
    rows = [[("", {}), ("Triton", {}), ("intj", {})]]
    for m, t_, i_ in zip(MACHINES, tr_host, ij_host):
        rows.append(
            [B(SHORT[m]), B(f2(t_), color=T.ink2), B(f"{i_:.3f}", color=T.intj)]
        )
    table(
        s,
        X,
        1.98,
        TW,
        rows,
        [1.85, (TW - 1.85) / 2, (TW - 1.85) / 2],
        row_h=0.48,
        size=12,
    )
    card(s, X, 4.15, TW, 1.45)
    text(
        s,
        X + 0.2,
        4.15,
        TW - 0.4,
        1.45,
        [
            [B("Triton: "), "grid=(0,) launch, skips only the driver."],
            [B("intj: "), "no_gpu launcher, driver call compiled out."],
            [B("Driver: "), "~3 µs on HIP, ~1.2 µs on CUDA."],
        ],
        size=12,
        bullets=True,
        gap=8,
        color=T.ink2,
        anchor="m",
    )
    source(
        s,
        "Source: data/headline_launch.csv (chart); data/sweep_args_wide.csv (table: triton_host, intj_host). "
        "Triton 3.8.0, intj 05e25a9.",
    )


def results_sweeps(d):
    s = d.slide(
        "intj stays flat as arguments, constexprs and cached kernels grow",
        notes=(
            f"Host-only intj = a no_gpu launcher on CPU tensors: {ARGS_KEY} and cache lookup. Tensor args, host only: "
            "Triton 4.9 -> 129.2 us on MI308X, 3.5 -> 87.5 on MI350X, 4.1 -> 27.7 on RTX 4080 SUPER; intj 0.022 -> "
            "0.298, 0.014 -> 0.130, 0.018 -> 0.123 us. Cache sweep: one constexpr takes K values in a fixed random "
            "order, so K kernels are cached and hit at random."
        ),
    )
    legend(s, MX, 1.5, [("Triton JITFunction", T.triton), ("intj", T.intj)])
    text(
        s,
        7.2,
        1.5,
        W - MX - 7.2,
        0.32,
        "MI308X, per launch incl. driver (µs)",
        size=12,
        color=T.muted,
        align="r",
    )
    panels = [
        (
            "Tensor arguments",
            [r for r in sweep_args if r["arg_kind"] == "tensor"],
            "n_args",
            "count",
        ),
        ("Constexpr arguments (+4 tensors)", sweep_const, "n_constexpr", "count"),
        (
            "Cached kernels, hit at random",
            sweep_cache,
            "n_cached_kernels",
            "kernels in cache",
        ),
    ]
    gap = 0.3
    pw = (CW - 2 * gap) / 3
    for i, (title, rows_all, key, xt) in enumerate(panels):
        x = MX + i * (pw + gap)
        rows = [r for r in rows_all if r["machine"] == M308]
        cats = [r[key] for r in rows]
        tr = [float(r["triton_launch_us"]) for r in rows]
        ij = [float(r["intj_launch_us"]) for r in rows]
        text(s, x, 1.95, pw, 0.35, title, size=14, bold=True)
        line_chart(
            s,
            x - 0.1,
            2.3,
            pw + 0.1,
            3.15,
            cats,
            [("Triton", tr, T.triton), ("intj", ij, T.intj)],
            x_title=xt,
        )
        card(s, x, 5.6, pw, 1.05)
        h = [float(r["intj_host_us"]) for r in rows]
        text(
            s,
            x + 0.18,
            5.6,
            pw - 0.36,
            1.05,
            [
                [
                    ("Triton ", {"color": T.ink2}),
                    B(f"{f1(tr[0])} → {f1(tr[-1])} µs", color=T.ink2),
                ],
                [
                    ("intj ", {"color": T.ink2}),
                    B(f"{f1(ij[0])} → {f1(ij[-1])} µs", color=T.intj),
                    (
                        f"   host only {h[0]:.3f} → {h[-1]:.3f}",
                        {"color": T.muted, "size": 11},
                    ),
                ],
            ],
            size=13,
            anchor="m",
            gap=4,
        )
    source(
        s,
        "Source: data/sweep_args_wide.csv, sweep_constexpr_wide.csv, sweep_cache_wide.csv (MI350X and "
        "RTX 4080 SUPER rows have the same shape). Empty @triton.jit kernels.",
    )


BLOG = DATA / "tilelang_ffi_blog"
BLOG_URL = "tilelang.com/blog/2025/12/11/ffi.html"
BLOG_SOURCE = (
    f"Source: {BLOG_URL}, \u201cTileLang x TVM FFI: Reducing Host Overhead and Building an "
    "FFI\u2011Based Compiler\u201d (LeiWang1999, Dec 11, 2025)."
)


def blog_date(s):
    """The blog's date, highlighted, and its title, under the slide title."""
    shp = rect(s, MX, 1.42, 1.75, 0.42, T.intj, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    fill_frame(
        shp.text_frame,
        "Dec 11, 2025",
        size=15,
        bold=True,
        color=T.on_accent,
        align="c",
        anchor="m",
    )
    text(
        s,
        MX + 1.95,
        1.42,
        CW - 1.95,
        0.42,
        [
            [
                B("TileLang × TVM FFI", color=T.ink),
                (f"  ·  {BLOG_URL}", {"color": T.muted}),
            ]
        ],
        size=13,
        anchor="m",
    )


def tilelang_half_of_figure2():
    """Figure 2's left half: its "TileLang" label stacked on the TileLang code, the blank between cut."""
    with Image.open(BLOG / "figure2_matmul_triton_tilelang.png") as im:
        label = im.crop((40, 72, 1240, 182))
        code = im.crop((40, 494, 1240, 1240))
    out = Image.new("RGB", (code.width, label.height + 40 + code.height), "white")
    out.paste(label, (0, 0))
    out.paste(code, (0, label.height + 40))
    return out


def citing_tilelang(d):
    s = d.slide(
        "Prior work: TileLang moves the host code into generated C++",
        notes=(
            "Quotes are verbatim from the blog; '\u2026' marks a cut. The second quote condenses: \"Taking the matmul "
            "kernel as an example: after the SplitHostDevice pass, TileLang splits the IR into device and host parts. "
            "The host IR looks like this: [listing] Finally, via codegen, we generate the C++ host code, which can be "
            "retrieved with matmul_kernel.get_host_source()\". The figure is the TileLang half of the blog's Figure 2, "
            "unmodified apart from the crop (the Triton half is the usual matmul_kernel[grid](...) with strides and "
            "BLOCK constexprs). Images: ppt/data/tilelang_ffi_blog/."
        ),
    )
    blog_date(s)
    bottom = picture(s, MX, 2.1, 5.1, tilelang_half_of_figure2())
    text(
        s,
        MX,
        bottom + 0.06,
        5.1,
        0.3,
        "Figure 2 (TileLang half): TileLang vs Triton matmul kernel",
        size=11,
        italic=True,
        color=T.muted,
    )
    X = MX + 5.1 + 0.45
    QW = W - MX - X
    quotes = [
        [
            (
                "Every time the user accesses torch.Tensor.shape (or device, stride, \u2026) in Python, it adds around ",
                {},
            ),
            B("0.1\u20131 \u00b5s of overhead"),
            (", which is painful when kernels are very fast.", {}),
        ],
        [
            (
                "\u2026 after the SplitHostDevice pass, TileLang splits the IR into device and host parts. \u2026 Finally, "
                "via codegen, ",
                {},
            ),
            B("we generate the C++ host code"),
            (".", {}),
        ],
        [
            (
                "For a DSL, CPU overhead is mostly kernel host and cache; TileLang leaves very little headroom in these two "
                "parts. (In contrast, due to design choices, Triton is weaker on these two aspects. In our experience Triton "
                "compiles faster, but ",
                {},
            ),
            B("its runtime CPU overhead is higher"),
            (". ", {}),
            B("We won\u2019t benchmark here"),
            ("; please feel free to try.)", {}),
        ],
    ]
    y = 2.1
    for q, h in zip(quotes, (1.0, 0.95, 1.95)):
        text(
            s,
            X,
            y - 0.12,
            0.5,
            0.6,
            "\u201c",
            size=40,
            bold=True,
            color=T.intj,
            font=T.head_font,
        )
        text(s, X + 0.5, y, QW - 0.5, h, [q], size=14, italic=True, color=T.ink2)
        y += h + 0.25
    source(s, BLOG_SOURCE)


def citing_tilelang_figures(d):
    s = d.slide(
        "TVM FFI: TileLang host overhead 2.1–3.8× lower; CuTeDSL ~2 µs",
        notes=(
            "Figure 3 (TileLang on an H800): per-call CPU overhead and first compile time, TVM FFI host code vs the earlier "
            'Cython wrapper; the blog: "improving kernel overhead by 2.1\u20133.8\u00d7 (this time even includes kernel '
            'launch!)". Figure 4 is a post by Tianqi Chen (Nov 28) the blog cites: CuteDSL 4.3.1, "Major host overhead '
            'optimization (10-40\u00b5s down to a 2\u00b5s in hot loops)", powered by the apache tvm-ffi ABI. The blog '
            "does not benchmark Triton; this deck's Triton host path, 4 tensor arguments, is 5.97-13.78 us "
            "(data/sweep_args_wide.csv). Different GPUs and kernels, so not a like-for-like comparison."
        ),
    )
    blog_date(s)
    LW = 8.1
    bottom = picture(s, MX, 2.1, LW, BLOG / "figure3_ffi_cpu_overhead.png")
    text(
        s,
        MX,
        bottom + 0.06,
        LW,
        0.3,
        "Figure 3: CPU host overhead reduction with TVM FFI",
        size=11,
        italic=True,
        color=T.muted,
    )
    tr = [float(sweep_row(m, 4)["triton_host_us"]) for m in MACHINES]
    card(s, MX, bottom + 0.5, LW, 6.8 - bottom - 0.5)
    text(
        s,
        MX + 0.22,
        bottom + 0.5,
        LW - 0.44,
        6.8 - bottom - 0.5,
        [
            [
                B("TileLang: "),
                "Cython \u2192 TVM FFI host code, 2.1\u20133.8\u00d7 less CPU overhead per call, kernel launch included.",
            ],
            [
                B("CuTeDSL 4.3.1: "),
                "host overhead 10\u201340 \u00b5s \u2192 ~2 \u00b5s in hot loops, on the tvm-ffi ABI.",
            ],
            [
                B("Triton: "),
                f"not benchmarked there; its host path measures {min(tr):.1f}\u2013{max(tr):.1f} \u00b5s here (4 tensor args).",
            ],
        ],
        size=13,
        bullets=True,
        gap=8,
        color=T.ink2,
        anchor="m",
    )
    X = MX + LW + 0.35
    RW = W - MX - X
    bottom4 = picture(s, X, 2.1, RW, BLOG / "figure4_ffi_cutedsl_conclusion.png")
    text(
        s,
        X,
        bottom4 + 0.06,
        RW,
        0.3,
        "Figure 4: CuTeDSL results adopting tvm\u2011ffi",
        size=11,
        italic=True,
        color=T.muted,
    )
    source(s, BLOG_SOURCE + " Triton host path: data/sweep_args_wide.csv.")


def results_ffi(d):
    def sb(m, path):
        return float(
            next(
                r
                for r in same_binary
                if r["machine"] == m and r["case"] == f"{path} same binary"
            )["us_per_call"]
        )

    ratio = [sb(m, "Triton") / sb(m, "INTJ") for m in MACHINES]
    s = d.slide(
        f"Same binary: intj matches TVM FFI and beats Triton by {min(ratio):.1f}–{max(ratio):.1f}×",
        notes=(
            "Left: one Triton binary per machine of an empty kernel with one tensor argument (HSACO on the MI308X and "
            "MI350X, cubin on the RTX 4080 SUPER), launched by TVM FFI, Triton and intj. TVM FFI loads it with "
            "hipModuleLoadData on ROCm and with its own cubin launcher (TVM_FFI_EMBED_CUBIN) on CUDA, as in its "
            "examples/cubin_launcher; intj reuses Triton's loaded function. On the MI350X one of three rounds caught the "
            "low hipModuleLaunchKernel state (2.2-2.3 us on both sides). Right: ffi_sweep, empty kernels with N "
            "arguments, where TVM FFI and intj launch different binaries; at 0 arguments FFI's kernel takes a cheaper "
            "stream path. TVM FFI gets preconverted TensorViews, intj takes torch tensors. apache-tvm-ffi 0.1.14.post1."
        ),
    )

    LW = 5.15
    text(
        s, MX, 1.55, LW, 0.35, "Same binary, 1 tensor argument (µs)", size=14, bold=True
    )
    legend(s, MX, 1.95, [("Triton", T.triton), ("TVM FFI", T.third), ("intj", T.intj)])
    bar_chart(
        s,
        MX - 0.1,
        2.35,
        LW + 0.1,
        3.65,
        [SHORT[m] for m in MACHINES],
        [
            ("Triton", [sb(m, "Triton") for m in MACHINES], T.triton),
            ("TVM FFI", [sb(m, "TVM FFI") for m in MACHINES], T.third),
            ("intj", [sb(m, "INTJ") for m in MACHINES], T.intj),
        ],
        fmt="0.00",
        gap=60,
        label_size=10,
    )

    X = MX + LW + 0.35
    RW = W - MX - X
    text(
        s,
        X,
        1.55,
        RW,
        0.35,
        "By argument count, different binaries (µs)",
        size=14,
        bold=True,
    )
    legend(s, X, 1.95, [("TVM FFI", T.third), ("intj", T.intj)])
    gap = 0.15
    pw = (RW - 2 * gap) / 3

    def pick(rows, ns, case):
        return [
            float(
                next(r for r in rows if r["n_args"] == n and r["case"] == case)[
                    "us_per_call"
                ]
            )
            for n in ns
        ]

    top = max(
        float(r["us_per_call"])
        for r in ffi_sweep
        if r["case"] in ("FFI empty kernel", "INTJ static_compile kernel")
    )
    for i, m in enumerate(MACHINES):
        x = X + i * (pw + gap)
        rows = [r for r in ffi_sweep if r["machine"] == m]
        ns = list(dict.fromkeys(r["n_args"] for r in rows))
        text(s, x, 2.35, pw, 0.3, SHORT[m], size=12, bold=True, align="c")
        line_chart(
            s,
            x - 0.05,
            2.65,
            pw + 0.1,
            3.35,
            ns,
            [
                ("TVM FFI", pick(rows, ns, "FFI empty kernel"), T.third),
                ("intj", pick(rows, ns, "INTJ static_compile kernel"), T.intj),
            ],
            x_title="arguments",
            y_max=float(int(top) + 1),
        )

    card(s, MX, 6.15, CW, 0.62)
    text(
        s,
        MX + 0.2,
        6.15,
        CW - 0.4,
        0.62,
        "intj takes torch tensors; TVM FFI gets preconverted TensorViews. At 0 arguments FFI's kernel takes a "
        "cheaper stream path.",
        size=12,
        color=T.ink2,
        anchor="m",
    )
    source(
        s,
        "Source: data/same_binary_1arg.csv (one HSACO / cubin per machine), data/ffi_sweep.csv (different "
        "binaries). Triton 3.8.0, intj 05e25a9.",
    )


def results_aiter(d):
    s = d.slide(
        "AITER ops: 2–4.5× less host time per call, outputs bitwise equal",
        notes=(
            "_FastLaunch is AITER's hand-written Python launch cache for FlashKDA; intj still halves its host time. "
            "Kernel times from rocprofv3 --kernel-trace are the same on both sides; outputs are bitwise equal on every "
            "shape. AITER 874d253dd3 (upstream) vs 5d710e670 (intj)."
        ),
    )

    def small(r):
        return float(r["T"]) <= 1024  # 2 kernels per call

    ops = [
        ("fused KDA decode", op_host(kda, "base_triton_jit"), op_host(kda, "intj")),
        (
            "FlashKDA (Gluon)",
            op_host(flash, "base_fastlaunch", small),
            op_host(flash, "intj", small),
        ),
        ("fp8_mqa_logits", op_host(fp8, "base_triton_jit"), op_host(fp8, "intj")),
    ]
    text(
        s,
        MX,
        1.55,
        7,
        0.35,
        "Host time per eager call (µs), best upstream path vs intj",
        size=14,
        bold=True,
    )
    legend(s, MX, 1.95, [("upstream", T.triton), ("intj", T.intj)])
    bar_chart(
        s,
        MX - 0.1,
        2.35,
        7.2,
        4.45,
        [o[0] for o in ops],
        [
            ("upstream", [o[1] for o in ops], T.triton),
            ("intj", [o[2] for o in ops], T.intj),
        ],
        gap=70,
        label_size=12,
    )

    X = 8.05
    SW = W - MX - X
    flash_jit = op_host(flash, "base_triton_jit", small)

    def eager_range(op, n=None):
        v = [
            float(r["eager_speedup_vs_upstream"])
            for r in aiter_summary
            if r["op"] == op
            and r["route"] == "default"
            and r["eager_speedup_vs_upstream"]
        ]
        v = v[:n] if n else v
        return f"{min(v):.2f}–{max(v):.2f}×"

    up, ij = ops[0][1], ops[0][2]
    stat(
        s,
        X,
        1.55,
        SW,
        1.25,
        f"{f1(up / ij)}×",
        f"KDA decode host time, {f1(up)} → {f1(ij)} µs",
    )
    stat(
        s,
        X,
        2.95,
        SW,
        1.25,
        f"{f1(ops[1][1] / ops[1][2])}×",
        f"FlashKDA vs AITER's own _FastLaunch ({f1(flash_jit / ops[1][2])}× vs Triton JIT)",
    )
    stat(
        s,
        X,
        4.35,
        SW,
        1.25,
        f"{f1(ops[2][1] / ops[2][2])}×",
        f"fp8_mqa_logits host time, {f1(ops[2][1])} → {f1(ops[2][2])} µs",
    )
    text(
        s,
        X,
        5.8,
        SW,
        0.95,
        f"Eager time where upstream is host-bound: {eager_range('kda_decode', 4)} KDA decode (batch ≤ 16), "
        f"{eager_range('fp8_mqa_logits')} fp8_mqa_logits. GPU kernels unchanged.",
        size=12,
        color=T.ink2,
    )
    source(
        s,
        "Source: data/aiter_*.csv, data/aiter_summary.csv, MI350X. FlashKDA: T ≤ 1024, 2 kernels per call. "
        "Medians over shapes.",
    )


def results_aiter_crossover(d):
    # (op, rows, x key, x values, summary shape, x label); eager time is about max(host, GPU)
    kda_rows = [r for r in kda if float(r["batch"]) <= 64]
    fp8_rows = [
        r for r in fp8 if r["seq_q"] == r["seq_kv"] and float(r["seq_q"]) <= 2048
    ]
    graphs = []
    for op, rows, key, shape, label in [
        (
            "kda_decode",
            kda_rows,
            "batch",
            lambda r: f"{r['batch']}x12",
            "Fused KDA decode, H = 12: eager µs per call vs batch",
        ),
        (
            "fp8_mqa_logits",
            fp8_rows,
            "seq_q",
            lambda r: f"{r['seq_q']}x{r['seq_kv']}x64",
            "fp8_mqa_logits, H = 64: eager µs per call vs q = kv",
        ),
    ]:
        xs = list(dict.fromkeys(r[key] for r in rows))

        def val(p, rows=rows, key=key, xs=xs):
            return [
                float(
                    next(r for r in rows if r[key] == x and r["launch_path"] == p)[
                        "eager_us"
                    ]
                )
                for x in xs
            ]

        summary = {
            r["shape"]: r
            for r in aiter_summary
            if r["op"] == op and r["route"] == "default"
        }
        firsts = {x: next(r for r in rows if r[key] == x) for x in xs}
        gpu = [float(summary[shape(firsts[x])]["gpu_kernel_us"]) for x in xs]
        host_bound = [
            summary[shape(firsts[x])]["regime"] == "upstream host-bound" for x in xs
        ]
        graphs.append(
            (op, label, xs, val("base_triton_jit"), val("intj"), gpu, host_bound)
        )

    floor = [u for _, _, _, up, _, _, hb in graphs for u, h in zip(up, hb) if h]

    def step(name):  # intj tree, median over q = kv 128-1024
        return statistics.median(
            float(r["host_us"])
            for r in fp8_breakdown
            if r["side"] == "intj" and r["step"] == name
        )

    launch = step("kernel launch")
    outside = [
        ("−inf fill", step("torch.full/empty")),
        ("config", step("gfx950 heuristics")),
        ("slice", step("slice [:, :kv]")),
        ("free", step("free previous logits")),
        (
            "checks",
            sum(
                step(k)
                for k in (
                    "shape+asserts",
                    "Q.device",
                    "4x stride()",
                    "other dict+grid",
                    "current_device_stream",
                )
            ),
        ),
    ]
    breakdown_note = ", ".join(
        f"{k} {v:.2f} us" for k, v in outside + [("launch", launch)]
    )
    proto_x = [r["seq_q"] for r in fp8_proto]
    s = d.slide(
        f"KDA decode, fp8_mqa_logits: upstream floors at {round(min(floor))}–{round(max(floor))} µs",
        notes=(
            "MI350X. Eager time: 32 back-to-back calls timed by GPU events, about max(host, GPU). Host-bound means the "
            "upstream call's host time exceeds its GPU kernel time (data/aiter_summary.csv regime). Larger shapes stay "
            "GPU-bound on both sides: KDA decode batch 128 / 256 at 83 / 157 us, fp8_mqa_logits q = kv 4096 / 8192 at "
            "127 / 425 us. fp8_mqa_logits at q = kv <= 1024: intj sits at its call's own host time (~15 us), above the "
            "6-17 us kernels (the Gluon kernel plus torch's fill). Host time per call, intj tree, medians over q = kv "
            "128-1024 and 3 processes, timed per section of an instrumented copy of the wrapper checked bitwise-equal "
            f"to it: {breakdown_note}. The prototype is the same wrapper with the gfx950 config cached per shape and the "
            "logits slice skipped when seq_len_kv is 256-aligned (all but 128); clean_logits=True, outputs bitwise equal "
            "on every shape. Dropping the -inf fill too (clean_logits=False) measured ~5.5 us host; not shown, it "
            "changes the output contract."
        ),
    )
    legend(
        s,
        MX,
        1.48,
        [
            ("upstream (Triton JIT)", T.triton),
            ("intj", T.intj),
            ("intj + leaner wrapper (prototype)", T.intj, "dash"),
            ("GPU kernel time", T.third),
        ],
    )
    CWL = 7.45
    X = MX + CWL + 0.3
    RW = W - MX - X
    for i, (op, label, xs, up, ij, gpu, hb) in enumerate(graphs):
        y = 1.88 + i * 2.55
        text(s, MX, y, CWL, 0.3, label, size=12, bold=True)
        # GPU time first, so intj's line is drawn on top of it.
        series: list[tuple] = [
            ("GPU kernel", gpu, T.third),
            ("upstream", up, T.triton),
            ("intj", ij, T.intj),
        ]
        if op == "fp8_mqa_logits":
            proto = [
                float(
                    next(r for r in fp8_proto if r["seq_q"] == x)["prototype_eager_us"]
                )
                for x in xs
            ]
            assert set(xs) <= set(proto_x)
            series.append(("prototype", proto, T.intj, "dash"))
        line_chart(s, MX - 0.1, y + 0.28, CWL + 0.1, 2.2, xs, series, y_step=10)
        hb_up = [u for u, h in zip(up, hb) if h]
        hb_ij = [v for v, h in zip(ij, hb) if h]
        last_hb = xs[max(j for j, h in enumerate(hb) if h)]
        what = "batch" if op == "kda_decode" else "q = kv"
        lines = [
            [
                B("Upstream: "),
                f"flat {f1(min(hb_up))}–{f1(max(hb_up))} µs to {what} {last_hb}",
            ]
        ]
        if op == "kda_decode":
            lines.append(
                [
                    B("intj: "),
                    f"{f1(min(hb_ij))}–{f1(max(hb_ij))} µs, the kernel's own time",
                ]
            )
            top = header_card(
                s,
                X,
                y + 0.55,
                RW,
                1.6,
                "Fused KDA decode",
                rule=T.intj,
                head_h=0.45,
                size=14,
            )
            text(
                s,
                X + 0.2,
                top,
                RW - 0.4,
                1.1,
                lines,
                size=12,
                bullets=True,
                gap=8,
                color=T.ink2,
                anchor="m",
            )
            continue
        # where the prototype is host-bound itself: its eager time within 15% of its host time
        proto_rows = [next(r for r in fp8_proto if r["seq_q"] == x) for x in xs]
        proto_hb = [
            r
            for r in proto_rows
            if float(r["prototype_eager_us"]) <= 1.15 * float(r["prototype_host_us"])
        ]
        hb_proto = [float(r["prototype_eager_us"]) for r in proto_hb]
        lines += [
            [
                B("intj: "),
                f"{f1(min(hb_ij))}–{f1(max(hb_ij))} µs, its own host time; launch {launch:.1f} µs",
            ],
            [
                B("Outside intj: "),
                " · ".join(f"{k} {v:.1f}" for k, v in outside) + " µs",
            ],
            [
                B("Prototype: "),
                f"{f1(min(hb_proto))}–{f1(max(hb_proto))} µs to q = kv {proto_hb[-1]['seq_q']}; "
                "config cached, slice skipped",
            ],
        ]
        top = header_card(
            s, X, y + 0.28, RW, 2.2, "fp8_mqa_logits", rule=T.intj, head_h=0.45, size=14
        )
        text(
            s,
            X + 0.2,
            top,
            RW - 0.4,
            1.75,
            lines,
            size=11,
            bullets=True,
            gap=6,
            color=T.ink2,
            anchor="m",
        )
    source(
        s,
        "Source: data/aiter_kda_decode.csv, aiter_fp8_mqa_logits.csv, aiter_summary.csv (rocprofv3, regime), "
        "aiter_fp8_mqa_logits_breakdown.csv, aiter_fp8_mqa_logits_prototype.csv. MI350X.",
    )


def results_adoption(d):
    s = d.slide(
        "AITER moved to intj: 426 launchers, 11 launches left on Triton",
        notes=(
            "Diff 874d253dd3..5d710e670: 338 files, +15,670 / -7,271. Four make_launcher( calls declare 11 more launchers. "
            "144 FlyDSL .launch sites in 77 files are out of scope (different launch ABI). The -402 is a regex count "
            "(445 removed, 43 added)."
        ),
    )
    sw = (CW - 3 * 0.3) / 4
    for i, (big, label) in enumerate(
        [
            ("109", "commits, 2026-09-24 → 10-14"),
            ("426", "@make_launcher decorators (0 before)"),
            ("200", ".py files import intj"),
            ("−402", "net kernel[grid](…) launch lines"),
        ]
    ):
        stat(s, MX + i * (sw + 0.3), 1.55, sw, 1.3, big, label)

    text(s, MX, 3.15, 5.6, 0.35, "Still on Triton, by reason", size=14, bold=True)
    for i, (n, why) in enumerate(
        [
            ("3", "num_ctas > 1 (CGA / multicast)"),
            ("3", "Iris multi-GPU helpers"),
            ("3", "torch.compile graphs"),
            ("2", "PDL (launch_pdl)"),
        ]
    ):
        y = 3.62 + i * 0.72
        badge(s, MX, y, n, fill=T.mid, ink=T.on_mid, d=0.48)
        text(s, MX + 0.68, y, 5, 0.48, why, size=14, anchor="m")

    X = 6.6
    RW = W - MX - X
    top = header_card(
        s, X, 3.15, RW, 3.6, "How it was checked", rule=T.rule, head_h=0.45, size=14
    )
    text(
        s,
        X + 0.2,
        top + 0.15,
        RW - 0.4,
        2.9,
        [
            [B("Check sweep: "), "all 424 launchers check and build on intj 20e578f."],
            [
                B("Differential test: "),
                "247 kernels, ~7,300 launches vs Triton, bitwise equal (except atomic_add and two "
                "self-nondeterministic cases; at intj 429d8e1).",
            ],
            [
                B("gfx942 host time: "),
                "rope 68.8 → 19.1 µs, flash_attn_func fwd 102.7 → 45.1 µs, attn_res_gate 20.0 → 12.2 µs.",
            ],
            [B("Untested: "), "CUDA; gfx1250 is compile-identity only."],
        ],
        size=13,
        bullets=True,
        gap=7,
        color=T.ink2,
    )
    source(
        s,
        "Source: aiter 874d253dd3..5d710e670 (git rev-list / git grep), aiter docs/intj_launch_exceptions.md, "
        "MIGRATION_GAP.md.",
    )


def gain_price(s, y, h, gains, prices, size=13):
    cw = (CW - 0.3) / 2
    for i, (items, label, rule) in enumerate(
        [(gains, "Gain", T.intj), (prices, "Price", T.muted)]
    ):
        x = MX + i * (cw + 0.3)
        top = header_card(s, x, y, cw, h, label, rule=rule, head_h=0.45)
        text(
            s,
            x + 0.22,
            top + 0.15,
            cw - 0.44,
            h - 0.7,
            items,
            size=size,
            bullets=True,
            gap=7,
            color=T.ink2,
        )


def micro(m, case):
    return float(
        next(r for r in incompat if r["machine"] == m and r["case"] == case)[
            "ns_per_call"
        ]
    )


def impl_launch_syntax(d):
    wrap = [
        micro(m, "kernel[grid](...), run stubbed")
        - micro(m, "kernel.run(...) directly, stubbed")
        for m in MACHINES
    ]
    cd = [micro(m, "triton.cdiv") for m in MACHINES]
    s = d.slide(
        "Incompatible launch syntax: no per-call grid or kernel[grid]",
        notes=(
            "Micro-benchmarks (ppt/scripts/bench_incompatible.py): each statement inline in a timeit loop, 200k calls "
            "x 9 samples, minus the same loop around pass; median of 3 processes per machine. triton.cdiv is a "
            "ConstexprFunction: every host call runs an import, unwraps args and kwargs in comprehensions, and checks "
            "knobs; (x + y - 1) // y is its body, timed inline (as a plain def call: 52 / 32 / 52 ns). kernel[grid] returns a new lambda per call; calling it re-packs *args and **kwargs into "
            "run(grid=..., warmup=False, ...). To time only that layer, run is replaced by a stub on a fresh "
            "JITFunction, and the same arguments are passed to the stub directly. intj's grid_cpp: no_gpu launchers on "
            "CPU tensors (decode, key, lookup), grid passed in vs computed in C. Triton sites blocked by intj's refusals "
            "(MIGRATION_GAP.md): tuples ~41, Gluon layouts ~75, TensorWrapper 14; parameter defaults alone block 0."
        ),
    )
    code_box(
        s,
        MX,
        1.5,
        6.6,
        1.45,
        "# Triton: device and stream implicit; keywords, defaults\n"
        "kernel[(triton.cdiv(n, 128),)](x, y, o, n, BLOCK=128)\n"
        "# intj: device, stream, grid passed in; positional only\n"
        "k = make_launcher(kernel)\n"
        "k(dev, stream, (triton.cdiv(n, 128),), x, y, o, n, 128)",
        size=11,
    )
    code_box(
        s,
        MX + 6.6 + 0.3,
        1.5,
        CW - 6.6 - 0.3,
        1.45,
        "# Triton's kernel[grid] (triton/runtime/jit.py)\n"
        "def __getitem__(self, grid) -> T:\n"
        "    return lambda *args, **kwargs: self.run(\n"
        "        grid=grid, warmup=False, *args, **kwargs)",
        size=11,
    )
    LW = 5.75
    for y, h, label, rule, items in [
        (
            3.15,
            1.65,
            "Gain",
            T.intj,
            [
                "Nothing is queried per call: device, stream and grid come from the caller.",
                "grid_arg / grid_cpp compute the grid in C; grid_py is opt-in.",
            ],
        ),
        (
            4.95,
            1.9,
            "Price",
            T.muted,
            [
                "Every call site is rewritten; defaulted parameters are passed explicitly.",
                "Refused: TensorWrapper, tuples, Gluon layout objects, num_ctas > 1, scratch, PDL, launch hooks.",
                "Undeclared cache-invalidating knobs freeze at the first call.",
            ],
        ),
    ]:
        top = header_card(s, MX, y, LW, h, label, rule=rule, head_h=0.4, size=14)
        text(
            s,
            MX + 0.2,
            top + 0.1,
            LW - 0.4,
            h - 0.55,
            items,
            size=12,
            bullets=True,
            gap=5,
            color=T.ink2,
        )
    X = MX + LW + 0.35
    TW = W - MX - X
    text(s, X, 3.15, TW, 0.3, "Host ns per call", size=14, bold=True)
    rows = [[("", {})] + [(SHORT[m], {"align": "r"}) for m in MACHINES]]
    for label, vals, hi in [
        ("triton.cdiv(n, 128)", cd, False),
        (
            "(x + y - 1) // y",
            [micro(m, "inline (n + b - 1) // b") for m in MACHINES],
            False,
        ),
        (
            "intj grid_cpp: cdiv in C",
            [
                micro(m, "intj, grid_cpp in C") - micro(m, "intj, grid passed in")
                for m in MACHINES
            ],
            True,
        ),
        ("kernel[grid] alone", [micro(m, "kernel[grid]") for m in MACHINES], False),
        ("kernel[grid](...) over run(...)", wrap, False),
    ]:
        cells = [
            ("≈0" if abs(v) < 0.5 else f"{v:+.0f}") if hi else f"{v:.0f}" for v in vals
        ]
        rows.append(
            [(label, {})]
            + [
                (c, {"align": "r", "color": T.intj if hi else T.ink, "bold": hi})
                for c in cells
            ]
        )
    table(s, X, 3.55, TW, rows, [2.75] + [(TW - 2.75) / 3] * 3, row_h=0.42, size=12)
    text(
        s,
        X,
        6.15,
        TW,
        0.65,
        "kernel[grid] builds that lambda on every launch; calling it packs *args and "
        "**kwargs, then unpacks them into run(). intj's grid_cpp is the extra over passing "
        "the grid in.",
        size=11,
        color=T.muted,
    )
    source(
        s,
        "Source: data/incompatible_micro.csv (scripts/bench_incompatible.py, 3 processes per machine), "
        "README.md, docs/Usage.md, MIGRATION_GAP.md.",
    )


def impl_ptr_range(d):
    rng = [
        micro(m, "triton get_tensor_specialization(tensor)") / 1000
        for m in (M308, M350)
    ]

    def per_tensor(m, col):  # slope of the host-only sweep, 0 -> 64 tensor arguments
        rows = [
            r for r in sweep_args if r["machine"] == m and r["arg_kind"] == "tensor"
        ]
        return (float(rows[-1][col]) - float(rows[0][col])) / (
            int(rows[-1]["n_args"]) - int(rows[0]["n_args"])
        )

    s = d.slide(
        f"Incompatible ptr_range: {min(rng):.1f}–{max(rng):.1f} µs per tensor in Triton",
        notes=(
            "On AMD, Triton adds the S bit (storage < 2 GiB) to every tensor's specialization: its native specializer "
            "calls HIPBackend.get_tensor_specialization (Python) for each tensor argument of each launch, which checks "
            "knobs.amd.use_buffer_ops and calls is_within_2gb. Micro-benchmarks: ppt/scripts/bench_incompatible.py, "
            "GPU tensor of 4096 floats, median of 3 processes. Per tensor argument from the host-only sweep slope "
            "(0 -> 64 tensors): Triton 1.94 / 1.31 / 0.37 us on MI308X / MI350X / RTX 4080 SUPER, where CUDA has no S "
            "bit; intj reads data pointer, dtype, alignment and storage size in C for every tensor. The MI350X host had "
            "a load average of ~8 during its run."
        ),
    )
    code_box(
        s,
        MX,
        1.5,
        6.3,
        1.55,
        "# Triton, per tensor argument, per launch (HIPBackend)\n"
        'if hasattr(arg, "ptr_range"):  # duck-typed objects\n'
        "    return arg.ptr_range() <= 2**31-1\n"
        "if isinstance(arg, torch.Tensor):\n"
        "    return arg.untyped_storage().size() <= 2**31-1",
        size=11,
    )
    card(s, MX, 3.25, 6.3, 0.95)
    text(
        s,
        MX + 0.2,
        3.25,
        5.9,
        0.95,
        [
            [
                B("intj: ", color=T.intj),
                (
                    "the storage size is read through PyTorch's C++ structures, in C, as part of "
                    f"{ARGS_KEY}.",
                    {},
                ),
            ]
        ],
        size=13,
        anchor="m",
    )
    sw = (6.3 - 0.3) / 2
    tri = [per_tensor(m, "triton_host_us") for m in MACHINES]
    ij = [per_tensor(m, "intj_host_us") * 1000 for m in MACHINES]
    stat(
        s,
        MX,
        4.4,
        sw,
        2.4,
        f"{tri[0]:.2f} µs",
        f"Triton host time per tensor argument, MI308X ({tri[1]:.2f} MI350X; {tri[2]:.2f} RTX 4080 SUPER, "
        "which has no S bit)",
        color=T.triton if T.name == "light" else T.ink,
        big_size=26,
    )
    stat(
        s,
        MX + sw + 0.3,
        4.4,
        sw,
        2.4,
        f"{ij[0]:.1f} ns",
        f"intj host time per tensor argument, MI308X ({ij[1]:.1f} MI350X, {ij[2]:.1f} RTX 4080 SUPER), "
        "range included",
        big_size=26,
    )

    X = MX + 6.3 + 0.35
    RW = W - MX - X
    text(
        s,
        X,
        1.5,
        RW,
        0.3,
        "Triton's per-tensor check, MI308X, ns per call",
        size=14,
        bold=True,
    )
    cases = [
        ("get_tensor_specialization", "triton get_tensor_specialization(tensor)"),
        ("is_within_2gb", "triton is_within_2gb(tensor)"),
        ("untyped_storage().size()", "tensor.untyped_storage().size()"),
        ("import torch", "import torch"),
        ("hasattr(t, 'ptr_range')", "hasattr(tensor, 'ptr_range')"),
    ]
    rev = list(reversed(cases))  # a bar chart draws its first category at the bottom
    bar_chart(
        s,
        X - 0.1,
        1.85,
        RW + 0.1,
        3.9,
        [c[0] for c in rev],
        [
            ("MI308X", [micro(M308, c[1]) for c in rev], T.triton),
        ],
        horizontal=True,
        fmt="0",
        gap=50,
        label_size=11,
    )
    text(
        s,
        X,
        5.85,
        RW,
        0.95,
        f"MI350X: {micro(M350, cases[0][1]):.0f} / {micro(M350, cases[1][1]):.0f} / "
        f"{micro(M350, cases[2][1]):.0f} / {micro(M350, cases[3][1]):.0f} / "
        f"{micro(M350, cases[4][1]):.0f} ns. The rest of is_within_2gb is isinstance, a "
        "second hasattr and the call itself.",
        size=11,
        color=T.muted,
    )
    source(
        s,
        "Source: data/incompatible_micro.csv (scripts/bench_incompatible.py), data/sweep_args_wide.csv "
        "(per-tensor slopes), triton/backends/amd/compiler.py (Triton 3.8.0).",
    )


def impl_lazy(d):
    s = d.slide(
        "Lazy building: make_launcher records; the first call builds",
        notes=(
            "Each launcher is a PyCFunction over a PyMethodDef embedded in its launcher object; the first call swaps "
            "ml_meth from the build shim to the rendered entry. A cache miss costs ~300 ns more than before "
            "(benchmark 404.6 -> ~707 ns, ~+180 ns in production). docs/Usage.md quotes ~250 ms for a build; the "
            "README's 0.6-1.9 s is used here. The AITER example is condensed: the gfx950 options also set num_stages=1, "
            "and the import sits inside the function because arch_info queries the GPU when imported. Build lock: one thread builds, the others wait detached from the interpreter; a "
            "call from inside the launcher's own build raises RuntimeError instead of deadlocking (runtime/intj_lazy.c). "
            "ARM64: the build release-stores ml_meth, but CPython reads it with a plain load; x86-64 keeps the order, a "
            "free-threaded ARM64 build may not (TODO.md). intj runs on x86-64 only today."
        ),
    )
    LW = 6.1
    arrow_x = MX + 0.21
    conn = s.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(arrow_x),
        Inches(1.75),
        Inches(arrow_x),
        Inches(4.0),
    )
    conn.line.color.rgb = rgb(T.card2)
    conn.line.width = Pt(2)
    for i, (head, body) in enumerate(
        [
            (
                "@make_launcher",
                "Records the kernel and options: no GPU, no torch import.",
            ),
            (
                "launcher.__self__.check()",
                "Optional: runs delayed= and every GPU-free refusal now.",
            ),
            (
                "First call",
                "delayed= (if check() did not), render, compile, load, swap ml_meth: "
                "0.6–1.9 s cold, 8.5 ms with the .so on disk.",
            ),
            (
                "Every later call",
                "The native entry, even through references taken before the build.",
            ),
        ]
    ):
        y = 1.5 + i * 0.8
        badge(
            s,
            MX,
            y,
            str(i + 1),
            fill=T.intj if i == 3 else T.mid,
            ink=T.on_accent if i == 3 else T.on_mid,
        )
        text(
            s,
            MX + 0.65,
            y - 0.02,
            LW - 0.65,
            0.3,
            head,
            size=13,
            bold=True,
            font=CODE if i < 2 else T.font,
        )
        text(s, MX + 0.65, y + 0.3, LW - 0.65, 0.45, body, size=11, color=T.ink2)
    text(
        s,
        MX,
        4.72,
        LW,
        0.3,
        "AITER: an option that needs the GPU goes in delayed=",
        size=12,
        bold=True,
    )
    text(
        s,
        MX,
        5.0,
        LW,
        0.25,
        "_triton_kernels/attention/chunked_pa_prefill.py",
        size=10,
        font=CODE,
        color=T.muted,
    )
    code_box(
        s,
        MX,
        5.28,
        LW,
        1.57,
        "def _decode_launch_keywords():  # runs on the first call\n"
        "    from aiter.ops.triton.utils._triton import arch_info\n"
        '    if arch_info.get_arch() == "gfx950":  # queries the GPU\n'
        '        return {"options": {"num_warps": 1, "waves_per_eu": 2, ...}}\n'
        "    return {}\n"
        "\n"
        "@make_launcher(grid_arg=2, delayed=_decode_launch_keywords)",
        size=10,
    )
    X = MX + LW + 0.35
    RW = W - MX - X
    for y, h, label, rule, items in [
        (
            1.5,
            2.45,
            "Gain",
            T.intj,
            [
                "import intj loads neither triton nor torch; decorating needs no GPU.",
                "Hit path unchanged: 0.4–3.0 ns faster on every untuned row after the change.",
                "Swapping ml_meth beat a custom callable type by 1.4–5 ns per call.",
            ],
        ),
        (
            4.15,
            2.7,
            "Price",
            T.muted,
            [
                "First call pays the build: 0.6 s, or 1.9 s for STATIC_COMPILE.",
                "Refusals move to check() or the first call.",
                "Thread-safe first call: a build lock.",
                "ARM64: publishing the built entry to other threads is unverified (TODO).",
            ],
        ),
    ]:
        top = header_card(s, X, y, RW, h, label, rule=rule, head_h=0.45)
        text(
            s,
            X + 0.22,
            top + 0.12,
            RW - 0.44,
            h - 0.65,
            items,
            size=12,
            bullets=True,
            gap=6,
            color=T.ink2,
        )
    source(
        s,
        "Source: README.md (build times, MI300X), benchmarks/journals/2026-09-28_lazy-task5_0_8d95e66.md, "
        "docs/superpowers/specs/2026-10-13-delayed-launcher-options-design.md, commit 8d9d5d5.",
    )


def impl_autotune(d):
    s = d.slide(
        "Autotune and heuristics need a chain of hash tables, not one key",
        notes=(
            "No Triton tutorial or example reads an unkeyed argument together with a tuned value in a heuristic. "
            "Triton's matmul (ops/matmul.py) keys K, so its EVEN_K = K % (BLOCK_K * SPLIT_K) == 0 is fixed per level-0 "
            "record and needs one lookup (tests/test_tuning.py::test_matmul_needs_one_lookup); Triton's "
            "test_decorator.py reads src.data_ptr() but no tuned value, which keys at level 0. The example is intj's "
            "tests/test_tuned.py::test_two_level_chain, run on a GPU: a bench that prefers BLOCK=64, calls N=8 with "
            "strides 128, 130, 192, asserting outputs 641, 640, 641, one tuning run (bench calls [32, 64]) and a "
            "two-key chain. A miss at either level runs a private copy of Triton's autotune/heuristics chain; its "
            "Autotuner cache already holds N=8, so a level-1 miss compiles without re-tuning. Level-0 map: spec words + "
            "exact N, swappable backend; child maps: intj's own map, 2 words per computed key, 4 slots to start; "
            "non-final records on the heap. Tuned hit vs plain intj +6.2 ns (MI308X); child-map hit 2.46 ns (map-only); "
            "never-coarser test over 40 cases plus its mutation test. Not measured: a Triton @autotune hit, a full "
            "two-level lookup end to end."
        ),
    )
    CWID = 7.7
    code_box(
        s,
        MX,
        1.5,
        CWID,
        2.05,
        "# intj tests/test_tuned.py::test_two_level_chain, condensed\n"
        '@triton.autotune(configs=[Config({"BLOCK": 32}), Config({"BLOCK": 64})],\n'
        '                 key=["N"])\n'
        '@triton.heuristics({"ALIGNED": lambda a: a["stride"] % a["BLOCK"] == 0})\n'
        "@triton.jit\n"
        "def aligned_tag(x, out, N, stride, BLOCK: tl.constexpr, ALIGNED: tl.constexpr): ...\n"
        "\n"
        "launch = make_launcher(aligned_tag)\n"
        "for stride in (128, 130, 192):\n"
        "    launch(dev, stream, 1, x, out, 8, stride)          # N = 8",
        size=10,
    )
    X = MX + CWID + 0.3
    RW = W - MX - X
    text(
        s,
        X,
        1.5,
        RW,
        2.1,
        [
            [
                B("ALIGNED "),
                "needs BLOCK, which only tuning for N picks: it cannot be keyed before a lookup.",
            ],
            [
                B("Keying stride "),
                "itself would tune again per stride; Triton tunes once per N.",
            ],
            [B("So: "), "level 0 keyed by N holds BLOCK; level 1 keyed by ALIGNED."],
        ],
        size=12,
        bullets=True,
        gap=6,
        color=T.ink2,
    )

    hdr = lambda v: (v, {"align": "l"})  # noqa: E731
    hit, miss = T.intj, T.ink
    rows = [
        [
            hdr("Call, N = 8"),
            hdr("Level 0: cache[N]"),
            hdr("C computes ALIGNED"),
            hdr("Level 1: child[ALIGNED]"),
            hdr("Python"),
        ],
        [
            ("stride 128", {"font": CODE}),
            ("miss → Triton tunes: BLOCK = 64", {"color": miss}),
            ("128 % 64 == 0 → 1", {"font": CODE}),
            ("miss → compile (64, 1)", {"color": miss}),
            "tune + compile",
        ],
        [
            ("stride 130", {"font": CODE}),
            ("hit: BLOCK = 64", {"color": hit, "bold": True}),
            ("130 % 64 == 2 → 0", {"font": CODE}),
            ("miss → compile (64, 0)", {"color": miss}),
            "compile, no re-tune",
        ],
        [
            ("stride 192", {"font": CODE}),
            ("hit: BLOCK = 64", {"color": hit, "bold": True}),
            ("192 % 64 == 0 → 1", {"font": CODE}),
            ("hit → launch", {"color": hit, "bold": True}),
            ("none", {"color": hit, "bold": True}),
        ],
    ]
    table(s, MX, 3.8, CW, rows, [1.6, 3.0, 2.85, 2.6, CW - 10.05], row_h=0.44, size=12)

    sw = (CW - 0.6) / 3
    for i, (big, label) in enumerate(
        [
            ("+6.2 ns", "tuned hit vs plain intj, MI308X"),
            ("2.46 ns", "child-map hit, map-only benchmark"),
            ("40 cases", "never-coarser test, plus its mutation test"),
        ]
    ):
        stat(
            s,
            MX + i * (sw + 0.3),
            5.85,
            sw,
            0.95,
            big,
            label,
            big_size=18,
            label_size=11,
        )
    source(
        s,
        "Source: tests/test_tuned.py::test_two_level_chain, intj/tuning.py, "
        "docs/superpowers/specs/2026-09-26-autotune-heuristics-design.md, journal 2026-09-26_autotune-task5.",
    )


def impl_internals(d):
    s = d.slide(
        "Reading CPython, PyTorch and Triton internals, version-gated",
        notes=(
            "Default: STATIC_COMPILE when its toolchain is available, else RUNTIME_SHIM when a verified layout exists, "
            "else INTERPRETER; an explicitly chosen unavailable mode raises. torch_abi.toml rows are generated by "
            "abi_detect and cross-checked by cpp_detect, never hand-edited. Gap: cache-invalidating knobs beyond debug / "
            "instrumentation / fpsan are not keyed yet (TODO.md)."
        ),
    )
    mw = (CW - 0.6) / 3
    text(
        s,
        MX,
        1.5,
        CW,
        0.32,
        f"Torch access mode: {ARGS_KEY} per call, 3 tensors, MI308X",
        size=14,
        bold=True,
    )
    for i, (name, body, ns) in enumerate(
        [
            (
                "INTERPRETER",
                "pointer and storage size through CPython calls",
                dec(M308, "interpreter"),
            ),
            (
                "RUNTIME_SHIM",
                "torch's structs at offsets from a verified table",
                dec(M308, "runtime_shim"),
            ),
            (
                "STATIC_COMPILE",
                "compiled against torch's headers (the default)",
                dec(M308, "static_compile"),
            ),
        ]
    ):
        x = MX + i * (mw + 0.3)
        card(s, x, 1.9, mw, 1.2)
        text(s, x + 0.2, 2.0, mw - 1.6, 0.35, name, size=13, bold=True, font=CODE)
        text(
            s,
            x + mw - 1.45,
            1.95,
            1.25,
            0.45,
            f"{round(ns)} ns",
            size=20,
            bold=True,
            color=T.intj if i else T.muted,
            align="r",
        )
        text(s, x + 0.2, 2.45, mw - 0.4, 0.6, body, size=12, color=T.ink2)
    table(
        s,
        MX,
        3.3,
        CW,
        [
            ["Dependency", "What intj reads", "How it is guarded"],
            [
                "CPython int",
                "ob_digit (≤ 3.11); lv_tag, PyUnstable_Long_* (3.12+)",
                "PY_VERSION_HEX branches, 30-bit static_assert, boundary tests",
            ],
            [
                "CPython calls",
                "vectorcall; the ml_meth swap; sizeof(PyObject) 16 / 32 (free-threaded)",
                "allowlist 3.8–3.14, 3.13t, 3.14t; module keyed by exact version + build",
            ],
            [
                "Torch tensor",
                "cdata; TensorImpl storage, offset, numel, dtype",
                "13 measured layouts (2.2–2.14), checked from C++; offset check at load",
            ],
            [
                "Triton",
                "native_specialize_impl, canonicalize_dtype, 3 knobs",
                "≥ 3.7 gate; version + libtriton.so in the key; never-coarser tests",
            ],
        ],
        [1.75, 5.0, CW - 6.75],
        row_h=0.5,
        size=12,
    )
    text(
        s,
        MX,
        5.95,
        CW,
        0.75,
        [
            [
                B("Layouts drift: "),
                (
                    "cdata moved 8 → 0 at torch 2.10, storage_offset 152 → 144 at 2.13. An unknown version "
                    "is refused, never borrowed from a neighbour. Offsets are x86-64 only.",
                    {"color": T.ink2},
                ),
            ]
        ],
        size=13,
    )
    source(
        s,
        "Source: CONCETPS.md, intj/python_intf/, intj/torch_intf/torch_abi.toml, intj/_imports.py, TODO.md. "
        f"{ARGS_KEY} per call: data/headline_decode_by_access_mode.csv.",
    )


def impl_tests(d):
    s = d.slide(
        "Test matrix: 149 GPU-free CI jobs, the full suite on two GPUs",
        notes=(
            "Each CI job installs the wheel, pins torch, and runs test_runtime.py in all three access modes: 446 job x "
            "mode runs. torch 2.9.1 on 3.14t skips INTERPRETER. Python 3.8 / 3.9 cells build without triton, since every "
            "triton >= 3.7 needs Python >= 3.10. gpu.yml has 12 ROCm cells for a self-hosted gfx942 runner, untested "
            "until one is registered. 1,075 is pytest --collect-only, not a pass count."
        ),
    )
    sw = (CW - 0.9) / 4
    for i, (big, label) in enumerate(
        [
            ("15", "OS × Python cells on 6 base images"),
            ("13", "torch pins, 2.2 → 2.14"),
            ("149", "jobs, each in all 3 access modes"),
            ("1,075", "tests collected on gfx942"),
        ]
    ):
        stat(s, MX + i * (sw + 0.3), 1.55, sw, 1.2, big, label)
    table(
        s,
        MX,
        3.0,
        5.3,
        [
            ["Python", "Cells", "Torch pins"],
            ["3.8", "1", "2.2–2.4"],
            ["3.9", "2", "2.2–2.8"],
            ["3.10 / 3.11 / 3.12", "2 / 3 / 3", "2.2–2.14"],
            ["3.13", "1", "2.5–2.14"],
            ["3.14 / 3.14t", "2 / 1", "2.9–2.14"],
        ],
        [2.2, 1.3, 1.8],
        row_h=0.45,
        size=13,
    )
    text(
        s,
        MX,
        5.85,
        5.3,
        0.75,
        "Ubuntu 22.04 / 24.04 / 26.04, Debian 12, AlmaLinux 9, manylinux_2_28. "
        "3.13t: local script only.",
        size=12,
        color=T.muted,
    )
    X = 6.25
    RW = W - MX - X
    top = header_card(
        s, X, 3.0, RW, 3.75, "Testing without a GPU", rule=T.intj, head_h=0.45, size=14
    )
    text(
        s,
        X + 0.22,
        top + 0.15,
        RW - 0.44,
        3.0,
        [
            [
                B("Stub driver: "),
                "the real rendered module, built with cc, records grid, block, stream and params.",
            ],
            [
                B("no_gpu=True: "),
                f"no target query, no driver dlopen; only {ARGS_KEY} and the cache lookup run.",
            ],
            [
                B("Fake driver: "),
                "tests patch the driver to device 0, stream 0, a hip / host target.",
            ],
            [B("No-GPU-query proof: "), "the target query is patched to raise."],
            [
                B("On GPUs: "),
                "full suite on MI308X (triton 3.7.1, 3.8.0) and RTX 4080 SUPER (869 passed, 21 skipped at 9682f5c).",
            ],
        ],
        size=12,
        bullets=True,
        gap=6,
        color=T.ink2,
    )
    source(
        s,
        "Source: ci/matrix.json, ci/image.py, ci/AGENTS.md, tests/test_runtime.py, tests/test_compat.py, README.md.",
    )


def kmap_ns(words, op, entries, backend):
    return float(
        next(
            r
            for r in kmap
            if r["key_words"] == str(words)
            and r["op"] == op
            and r["entries"] == str(entries)
            and r["backend"] == backend
        )["ns"]
    )


def impl_micro(d):
    s = d.slide(
        "Micro optimization: the call, and a map built for 1–8 entries",
        notes=(
            "Map numbers: tests/bench_kernel_cache.cpp (google/benchmark, map only, no last-key memo, hash precomputed by "
            "the caller, uniformly random keys), from benchmarks/journals/2026-10-14_full-suite_0_bc8caf7.md (MI308X "
            "host, Xeon 8480C, 3 rounds); data/kernel_cache_map.csv. intj: open addressing, linear probing, values "
            "stored inline, 16 slots to start, doubling at 50% load, an empty slot is one whose value reads as zero. "
            "tsl = tsl::robin_map, absl = absl::flat_hash_map, the swappable kernel_cache backends. hit/1 always touches "
            "slot 0 (a permanently hot line). Real launchers hold 1-8 specializations. A typical key is 2 words (one "
            "word of per-parameter code bytes, one constexpr value)."
        ),
    )
    items = [
        (
            "1",
            "Specialized entry + FASTCALL",
            "183 → 118 ns",
            "32 int args, same key (c7c1be5)",
            [
                "METH_FASTCALL lets CPython specialize the call site to CALL_BUILTIN_FAST."
            ],
        ),
        (
            "2",
            "Hash table for a sparse load",
            "107 → 69 ns",
            "32 constexpr args, same key (7ad507d)",
            [
                "Open addressing, values inline, 50% max load; real launchers hold 1–8 entries.",
                "A one-entry memo of the last hit is compared before hashing.",
            ],
        ),
    ]
    cw = (CW - 0.3) / 2
    ch = 1.75
    for i, (n, head, big, what, bl) in enumerate(items):
        x = MX + i * (cw + 0.3)
        y = 1.5
        card(s, x, y, cw, ch)
        badge(s, x + 0.2, y + 0.2, n)
        text(s, x + 0.75, y + 0.2, cw - 3.3, 0.42, head, size=14, bold=True, anchor="m")
        text(
            s,
            x + cw - 2.6,
            y + 0.12,
            2.4,
            0.5,
            big,
            size=20,
            bold=True,
            color=T.intj,
            align="r",
            anchor="m",
        )
        text(
            s, x + cw - 3.7, y + 0.64, 3.5, 0.3, what, size=10, color=T.muted, align="r"
        )
        text(
            s,
            x + 0.25,
            y + 1.0,
            cw - 0.5,
            ch - 1.05,
            bl,
            size=12,
            bullets=True,
            gap=4,
            color=T.ink2,
        )

    backends = [
        ("intj", T.intj),
        ("tsl::robin_map", T.third),
        ("absl::flat_hash_map", T.mid),
    ]
    keys = ["intj", "tsl", "absl"]
    LW = 6.6
    text(
        s,
        MX,
        3.45,
        LW,
        0.3,
        "Map hit (ns) vs live entries, 2-word key",
        size=13,
        bold=True,
    )
    legend(s, MX, 3.75, backends)
    entries = [1, 8, 64, 512]
    bar_chart(
        s,
        MX - 0.1,
        4.05,
        LW + 0.1,
        2.2,
        [str(e) for e in entries],
        [
            (name, [kmap_ns(2, "hit", e, k) for e in entries], color)
            for (name, color), k in zip(backends, keys)
        ],
        fmt="0.0",
        gap=60,
        label_size=10,
    )
    X = MX + LW + 0.35
    TW = W - MX - X
    text(
        s, X, 3.45, TW, 0.3, "Hit at 8 entries (ns), by key length", size=13, bold=True
    )
    rows = [
        [
            ("", {}),
            ("intj", {"align": "r"}),
            ("tsl", {"align": "r"}),
            ("absl", {"align": "r"}),
        ]
    ]
    for words in (1, 2, 5):
        rows.append(
            [(f"{words}-word key", {})]
            + [
                (
                    f"{kmap_ns(words, 'hit', 8, k):.2f}",
                    {
                        "align": "r",
                        "bold": k == "intj",
                        "color": T.intj if k == "intj" else T.ink,
                    },
                )
                for k in keys
            ]
        )
    table(s, X, 3.85, TW, rows, [1.55] + [(TW - 1.55) / 3] * 3, row_h=0.42, size=12)
    text(
        s,
        X,
        5.6,
        TW,
        0.6,
        "Same keys, hash precomputed, no last-key memo; MI308X host (Xeon 8480C).",
        size=11,
        color=T.muted,
    )

    rect(s, MX, 6.35, CW, 0.47, T.card2)
    text(
        s,
        MX + 0.2,
        6.35,
        CW - 0.4,
        0.47,
        [
            [
                B("The loop: ", color=T.intj),
                (
                    "bench_launch.py --no-gpu | --sweep | --last-key on a pinned core; 5 "
                    "interleaved processes; a row regresses at > 3% and > 2 ns.",
                    {},
                ),
            ]
        ],
        size=12,
        anchor="m",
    )
    source(
        s,
        "Source: benchmarks/journals/2026-09-28_mix-history_0_6e0d4bc.md (1, 2), data/kernel_cache_map.csv "
        "(tests/bench_kernel_cache.cpp, journal 2026-10-14_full-suite_0_bc8caf7.md).",
    )


def disp(m, case):
    return float(
        next(r for r in ffi_disp if r["machine"] == m and r["case"] == case)["value"]
    )


def impl_ffi_dispatch(d):
    per_int = {
        m: (disp(m, "light: 4 ints") - disp(m, "light: 0 args")) / 4 for m in MACHINES
    }
    over = {
        m: disp(m, "unordered_map find, light") - disp(m, "direct type compare")
        for m in MACHINES
    }
    s = d.slide(
        "TVM FFI: one unordered_map lookup per argument",
        notes=(
            "TVM FFI apache-tvm-ffi 0.1.14.post1 (PyPI wheel). Every Python argument goes through "
            "TVMFFIPyCallManager::SetArgument: Py_TYPE(arg), then std::unordered_map<PyTypeObject*, setter>::find "
            "(thread-local, reserve(32)); a type seen first goes through TVMFFICyArgSetterFactory and is emplaced. "
            "On the C++ side, Any -> T is TypeTraits<T>::TryCastFromAnyView: a type-index switch, no lookup. "
            "Isolated: ppt/scripts/bench_unordered_map_lookup.cpp, built with each machine's g++ -O2 (13.3 / 15.2 / "
            "16.2 -- the wheel's own build is not this code, so read the isolated rows as the shape of the cost), keys "
            "are PyTypeObject-sized heap blocks, 4 lookups per call, 20M x 9 samples, median of 3 processes; light = 6 "
            "entries, heavy = 1,030 (1,493 buckets); spread keys (4,096 draws over all entries) read 4.3 / 2.4-2.5 / "
            "1.4-1.5 ns. End to end: ppt/scripts/bench_ffi_dispatch.py, testing.nop inline in timeit, light = the 2 "
            "types this run uses, heavy = after 1,024 distinct int subclasses. A heavily loaded map costs the same in "
            "steady state: the bucket and node of a hot type stay cached. 4 CPU tensors through TVM FFI: 541 / 332 / "
            "327 ns per call, ~110 / ~68 / ~65 ns per tensor (DLPack exchange); intj's per-tensor host cost is "
            "1.6-4.3 ns (sweep slopes)."
        ),
    )
    LW = 6.45
    text(
        s,
        MX,
        1.5,
        LW,
        0.3,
        "PyObject* → TVMFFIAny: a hash lookup per argument",
        size=13,
        bold=True,
    )
    code_box(
        s,
        MX,
        1.82,
        LW,
        2.15,
        "// python/tvm_ffi/cython/tvm_ffi_python_helpers.h, SetArgument\n"
        "PyTypeObject* py_type = Py_TYPE(py_arg);\n"
        "auto it = arg_dispatch_map_.find(py_type);  // unordered_map\n"
        "if (TVM_FFI_PREDICT_TRUE(it != arg_dispatch_map_.end())) {\n"
        "  TVMFFIPyArgSetter setter = it->second;\n"
        "  // if error happens, propagate it back\n"
        "  if (TVM_FFI_PREDICT_FALSE(setter(ctx, py_arg, out) != 0)) return -1;\n"
        "} ...",
        size=10,
    )
    text(
        s,
        MX,
        4.15,
        LW,
        0.3,
        "TVMFFIAny → DLTensor / int: a type-driven cast, no lookup",
        size=13,
        bold=True,
    )
    code_box(
        s,
        MX,
        4.47,
        LW,
        0.75,
        "// include/tvm/ffi/any.h\n"
        "std::optional<T> opt = TypeTraits<T>::TryCastFromAnyView(&data_);",
        size=10,
    )
    card(s, MX, 5.42, LW, 1.4)
    text(
        s,
        MX + 0.2,
        5.42,
        LW - 0.4,
        1.4,
        [
            [
                B("Lookup cost: ", color=T.intj),
                f"{min(over.values()):.1f}–{max(over.values()):.1f} ns per argument over a "
                "direct type compare, about a tenth of TVM FFI's "
                f"{min(per_int.values()):.0f}–{max(per_int.values()):.0f} ns per int argument.",
            ],
            [
                B("Heavy load: ", color=T.intj),
                "1,000+ entries cost the same as 2 in steady state.",
            ],
            [
                B("intj: ", color=T.intj),
                f"each parameter's check is compiled in; {ARGS_KEY} has no lookup.",
            ],
        ],
        size=12,
        bullets=True,
        gap=4,
        color=T.ink2,
        anchor="m",
    )

    X = MX + LW + 0.35
    TW = W - MX - X
    head = [("", {})] + [(SHORT[m], {"align": "r"}) for m in MACHINES]
    text(
        s,
        X,
        1.5,
        TW,
        0.3,
        "Isolated, ns per argument (4-argument call)",
        size=13,
        bold=True,
    )
    rows = [head]
    for label, case, hi in [
        ("find, light (6 entries)", "unordered_map find, light", False),
        ("find, heavy (1,030 entries)", "unordered_map find, heavy", False),
        ("direct type compare", "direct type compare", True),
    ]:
        rows.append(
            [(label, {})]
            + [
                (
                    f"{disp(m, case):.2f}",
                    {"align": "r", "bold": hi, "color": T.intj if hi else T.ink},
                )
                for m in MACHINES
            ]
        )
    table(s, X, 1.85, TW, rows, [2.35] + [(TW - 2.35) / 3] * 3, row_h=0.42, size=11)
    text(
        s,
        X,
        3.7,
        TW,
        0.3,
        "testing.nop, ns per call (map: 2 / 1,026 types)",
        size=13,
        bold=True,
    )
    rows = [head]
    for label, case in [
        ("nop()", "light: 0 args"),
        ("nop(i, i, i, i), light", "light: 4 ints"),
        ("nop(i, i, i, i), heavy", "heavy: 4 ints"),
    ]:
        rows.append(
            [(label, {"font": CODE, "size": 10})]
            + [(f"{disp(m, case):.0f}", {"align": "r"}) for m in MACHINES]
        )
    rows.append(
        [("per int argument", {"bold": True})]
        + [(f"{per_int[m]:.1f}", {"align": "r", "bold": True}) for m in MACHINES]
    )
    table(s, X, 4.05, TW, rows, [2.35] + [(TW - 2.35) / 3] * 3, row_h=0.42, size=11)
    text(
        s,
        X,
        6.25,
        TW,
        0.55,
        "Median of 3 processes per machine; isolated rows built with the machine's g++.",
        size=10,
        color=T.muted,
    )
    source(
        s,
        "Source: data/ffi_dispatch.csv (scripts/bench_unordered_map_lookup.cpp, scripts/bench_ffi_dispatch.py), "
        "apache-tvm-ffi 0.1.14.post1 tvm_ffi_python_helpers.h, include/tvm/ffi/any.h.",
    )


def impl_annotation(d):
    s = d.slide(
        "Annotation: state facts, compile exactly those",
        notes=(
            "Argument(type, specialize, value, bind_value) and Constexpr(type, power_of_two_or_zero, value), inline on the "
            "kernel or via extra_annotation. verify_annotation=True costs 41.7-43.1 ns in the same benchmark. aiter "
            "23bf3e29b dropped the false Aligned(16) on stride_scale_seq; the HSACO was unchanged. Exactly arrived in "
            "intj 20e578f. power_of_two_or_zero (docs/Usage.md): an integer constexpr promised to be 0 or +-2**k is keyed "
            "in one byte (0 -> 0, 2**k -> k+1, -2**k -> 0x80 | (k+1)) instead of one 8-byte value word; "
            "verify_annotation=True rejects other values, the unchecked default trusts the promise. No host-time "
            "measurement of it on its own."
        ),
    )
    code_box(
        s,
        MX,
        1.5,
        6.3,
        2.95,
        "launch = make_launcher(kernel, extra_annotation={\n"
        '  "x": Argument(type=tl.pointer_type(tl.float32),\n'
        "                specialize=NEVER,\n"
        "                bind_value=BindValue.TENSOR),\n"
        '  "BLOCK": Constexpr(type=tl.int64,\n'
        "                     power_of_two_or_zero=True)})\n"
        "bound = launch.__self__.bind(x=tensor)\n"
        "bound(device, stream, grid, *rest)\n\n"
        "# AITER pa_mqa_logits\n"
        "kv = Exactly(Aligned(16)) if padded else NEVER",
        size=12,
    )
    mw = (6.3 - 0.3) / 2
    for i, (m, desc) in enumerate(
        [
            ("AUTO", "measured per call, keyed"),
            ("NEVER", "not specialized"),
            ("Assume", "listed facts promised, rest AUTO"),
            ("Exactly", "listed facts promised, rest NEVER"),
        ]
    ):
        x = MX + (i % 2) * (mw + 0.3)
        y = 4.65 + (i // 2) * 1.0
        card(s, x, y, mw, 0.85)
        text(
            s,
            x + 0.18,
            y + 0.1,
            mw - 0.36,
            0.3,
            m,
            size=13,
            bold=True,
            font=CODE,
            color=T.intj,
        )
        text(s, x + 0.18, y + 0.45, mw - 0.36, 0.35, desc, size=12, color=T.ink2)

    X = 7.25
    RW = W - MX - X
    text(
        s,
        X,
        1.55,
        RW,
        3.3,
        [
            [
                B("Facts: "),
                "EqualTo(1), Aligned(16), PointerRange(32); a torch.Tensor annotation makes the argument tensor-only.",
            ],
            [
                B("Smaller key: "),
                "Exactly and NEVER add 0 key bytes; baked and bound values leave the call.",
            ],
            [
                B("Power of two or zero: "),
                "Constexpr(power_of_two_or_zero=True) keys a BLOCK-style constexpr in "
                "1 byte, not an 8-byte word: 0 → 0, 2**k → k+1, −2**k → 0x80 | (k+1).",
            ],
            [
                B("Exact binaries: "),
                "AITER's pa_mqa_logits needs 6 Triton compiles instead of 15, identical HSACO on "
                "gfx942 and gfx950.",
            ],
            [
                B("Price, trust: "),
                "a false promise selects the wrong binary unless verify_annotation=True. Upstream's "
                "Aligned(16) on a stride of 36 was one.",
            ],
        ],
        size=13,
        bullets=True,
        gap=8,
        color=T.ink2,
    )
    text(s, X, 5.0, RW, 0.3, "Host-only ns per call, MI308X", size=12, bold=True)
    table(
        s,
        X,
        5.35,
        RW,
        [
            [
                ("none", {"align": "c"}),
                ("baked", {"align": "c"}),
                ("tensor-annotated", {"align": "c"}),
                ("typed int", {"align": "c"}),
            ],
            [
                ("39.8", {"align": "c"}),
                ("37.5", {"align": "c", "color": T.intj, "bold": True}),
                ("39.4", {"align": "c"}),
                ("43.7", {"align": "c"}),
            ],
        ],
        [1.05, 1.05, 1.75, RW - 3.85],
        row_h=0.42,
        size=12,
    )
    text(
        s,
        X,
        6.3,
        RW,
        0.45,
        "Host time is about neutral; the win is the key and the binaries.",
        size=12,
        italic=True,
        color=T.muted,
    )
    source(
        s,
        "Source: docs/Usage.md, intj/annotation.py, tests/test_annotation.py, "
        "benchmarks/journals/2026-10-14_full-suite_0_bc8caf7.md, aiter docs/intj_launch_exceptions.md.",
    )


def discussion(d):
    s = d.slide(
        "Discussion: next steps",
        notes=(
            "1. The AITER migration: branch codex/intj-all-launches on lijinpei-amd/aiter (head 5d710e670 measured; "
            "23bf3e29b is one commit later), base 874d253dd3 on ROCm/aiter main; the PR number was not looked up "
            "(no token access to the ROCm org here). Numbers from the AITER adoption slide. 2. Line counts: "
            "git ls-files intj | xargs wc -l, without torch_abi.toml: 12,488 (8,028 Python, 4,147 C / header / Jinja); "
            "git ls-files tests | xargs wc -l: 12,886. Most of it was written by coding agents."
        ),
    )
    cw = (CW - 0.4) / 2
    ch = 3.3
    for i, (head, items) in enumerate(
        [
            (
                "Merge the AITER PR",
                [
                    "Branch codex/intj-all-launches onto ROCm/aiter main: 109 commits, 426 launchers across 200 files.",
                    "247 kernels checked bitwise equal against Triton.",
                    "11 launches stay on Triton: num_ctas > 1, Iris, torch.compile, PDL.",
                ],
            ),
            (
                "Better docs and source code",
                [
                    "intj is 12.5K lines of source plus 12.9K of tests, mostly agent-generated.",
                    "Docs a newcomer can follow from first launch to annotation.",
                    "Source a human reviewer can read and own.",
                ],
            ),
        ]
    ):
        x = MX + i * (cw + 0.4)
        y = 1.6
        top = header_card(
            s, x, y, cw, ch, "", rule=T.intj if i == 0 else T.rule, head_h=0.75
        )
        badge(s, x + 0.25, y + 0.17, str(i + 1), d=0.5, size=16)
        text(
            s, x + 0.95, y + 0.05, cw - 1.2, 0.75, head, size=20, bold=True, anchor="m"
        )
        text(
            s,
            x + 0.3,
            top + 0.3,
            cw - 0.6,
            ch - 1.2,
            items,
            size=15,
            bullets=True,
            gap=10,
            color=T.ink2,
        )
    source(
        s,
        "Source: AITER adoption slide (aiter 874d253dd3..5d710e670); git ls-files intj / tests | xargs wc -l.",
    )


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--theme", choices=["amd", "light", "all"], default="all")
    ap.add_argument(
        "--template", default=AMD.template, help="the .potx for the amd theme"
    )
    args = ap.parse_args()
    themes = {"amd": [AMD], "light": [LIGHT], "all": [AMD, LIGHT]}[args.theme]
    for theme in themes:
        if theme.template:
            theme = replace(theme, template=args.template)
        out = ROOT / "results" / f"intj_launch_overhead_{theme.name}.pptx"
        out.parent.mkdir(parents=True, exist_ok=True)
        build(theme).save(str(out))
        print("wrote", out)


if __name__ == "__main__":
    main()
